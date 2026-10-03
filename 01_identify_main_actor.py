#!/usr/bin/env python3
# Both Sides project
# Requirements:
#   pip install openai openpyxl
#   export OPENAI_API_KEY="..."

import json
import os
import sys
from pathlib import Path
from openai import OpenAI
from openpyxl import load_workbook

MODEL = "gpt-5.6-terra"
REASONING_EFFORT = "medium"
CHUNK_SIZE = 250
MAX_ACTIVE_BATCHES = 2

STATE_FILE = "main_actor_state.json"
BATCH_DIR = "main_actor_batches"

INSTRUCTIONS = r"""
You are identifying the MAIN ACTOR responsible for the central political action described in an event.

You receive:
EVENT DESCRIPTION — the original event description.

Your task is to assign exactly one MAIN ACTOR from the allowed list.

ALLOWED MAIN ACTORS

1. settlers
2. protesters against government
3. arab israelis
4. haredi jews
5. right wing protesters
6. peace movements
7. not relevant


CORE MAIN-ACTOR RULE

The MAIN ACTOR is the group that INITIATES or PERFORMS the central action of the event.

Do NOT classify according to:
- the first group mentioned
- the group most frequently mentioned
- the victim
- the group being protested against
- police or military responding to another actor
- a secondary participant whose action is not the central action

If several actors appear, identify who performs the central action described in the record.


ACTOR DEFINITIONS


1. settlers

Israeli settlers, settler activists, settler groups, or settlement movements acting as the central actor.

This includes, when explicitly attributed to them:
- settler demonstrations
- attacks by settlers
- establishment or expansion of outposts
- land seizure or cultivation by settlers
- settler marches
- settler road blockades
- settler actions against Palestinians or Israeli authorities

Do not infer "settlers" merely from a West Bank location.
The acting group must be identified as settlers or clearly described as a settler group.


2. protesters against government

Use this umbrella category for Israeli protesters whose central political action is directed against the government or its policies, including:

- anti-judicial-reform / anti-judicial-overhaul demonstrators
- anti-government protesters
- protests calling for the government's resignation or elections
- hostage-deal / hostage-release protesters when the action is directed at the government or demands government action
- Kaplan protest movement and closely related protest groups
- protests combining these causes

This category deliberately combines the former labels:
- anti judicial reform demonstrators
- anti government protesters
- hostage deal protesters

Do not use it merely because someone criticizes a specific policy.
There must be a recognizable anti-government, judicial-overhaul, elections, or hostage-deal protest context.


3. arab israelis

Arab citizens or residents of Israel acting collectively in an Israeli political or civic context.

Use only when the description establishes that the central actors are Arab Israelis / Israeli Arabs / Arab citizens or residents of Israel.

Do NOT use for:
- Palestinians in the West Bank
- Palestinians in Gaza
- Palestinian residents when Israeli citizenship/residency context is not established
- a location that is Arab-majority without evidence about who acted

If identity is ambiguous, prefer not relevant rather than infer.


4. haredi jews

Haredi / ultra-Orthodox Jews acting collectively as the central actor.

Examples include:
- demonstrations against military conscription
- religious protests
- Haredi road blockades
- Haredi clashes or political demonstrations

Do not use merely because an event occurred in a Haredi neighborhood.


5. right wing protesters

Israeli right-wing / nationalist protesters acting collectively in a political demonstration or direct action.

Examples include:
- pro-judicial-reform demonstrations
- nationalist/right-wing demonstrations
- protests opposing government concessions from the right
- right-wing road blockades
- protests identified explicitly with right-wing organizations or causes

Do not classify settlers here when settlers themselves are the central actor.
Use settlers for explicitly settler-led actions.


6. peace movements

Israeli peace, anti-occupation, coexistence, or similar organized left-wing movements acting as the central actor.

Examples include explicitly identified:
- peace organizations
- anti-occupation activists
- coexistence movements
- Israeli groups protesting settlement activity or occupation

Do not use this category simply because a protest is left-wing.
Anti-government / judicial-overhaul / hostage-deal protests belong under protesters against government when that is their central context.


7. not relevant

Use not relevant when the central actor is outside the six project actor categories.

This includes:
- Palestinian actors in the West Bank or Gaza
- security forces acting alone
- police or military operations
- ordinary crime
- personal disputes
- accidents
- labor disputes or local protests unrelated to the project categories
- foreign actors
- events where the relevant actor is only a victim or secondary participant
- events with insufficient evidence to identify one of the allowed project actors


SECURITY-FORCES RULE

Police, IDF, Border Police, or other security forces do NOT become the MAIN ACTOR merely because they:
- disperse a protest
- arrest protesters
- use force after a protest action
- respond to stones, roadblocks, riots, or attacks

Classify according to the actor responsible for the central initiating action.

If the description is centrally about an independent police/military/security-force action, return not relevant.


MULTIPLE-ACTOR RULE

If several project actors are present:
- identify which group performed the central action
- do not merge categories
- do not choose a group merely because its action was more severe
- if the description genuinely does not establish a central actor, return not relevant with low certainty and needs_review = yes


FINAL CHECK

Before answering, verify:

1. Who actually performed the central action?
2. Did I accidentally classify the victim or responder?
3. Did I infer identity from location alone?
4. Does the actor fit one of the six project categories?
5. If the evidence is genuinely ambiguous, did I mark it for review?



SUPPLEMENTARY RULES (from the 2026-10-03 hand review — these override the definitions above where they conflict)

RIGHT WING PROTESTERS by demand, not by label:
A protest whose demand is to CONTINUE or RESUME a war, to REJECT a ceasefire or a hostage
deal, to BLOCK humanitarian aid, or to defend soldiers / "Jewish prisoners" against the
state is right wing protesters — even when the text only says "residents", "reservists",
"mothers of soldiers", "injured soldiers", "bereaved families", "Israelis". The demand
defines the camp. "Family members of hostages" taking part in an anti-aid or anti-deal
protest does not make it peace movements.

JEWISH NATIONALIST VIOLENCE OUTSIDE THE SETTLEMENTS:
Jews attacking Arabs inside Israel or in Jerusalem neighbourhoods (teens beating an Arab
bus driver, football fans attacking Arab drivers, a bus riot chanting "death to Arabs")
are right wing protesters, not settlers. Settlers is for the settlement movement acting
in the West Bank — including Israelis from inside Israel who raid a West Bank village,
and including the settlement movement's own activism wherever it happens (Nachala,
Hilltop Youth, Gaza/Lebanon "resettlement" marches and crossings, protests to free
detained settlers).

PROTESTERS AGAINST GOVERNMENT, specific cases:
- October Council / demands for a state commission of inquiry into 7 October.
- Protests against the war with Iran (2026).
- Brothers and Sisters in Arms / reservist protests against the government (draft
  exemption, corruption) that are not pro-war.
- Protests defending the Attorney General, the courts, or against a minister's conduct.

PEACE MOVEMENTS starts on 7 October 2023:
Hostage-deal, ceasefire, anti-war and anti-occupation protests from 7 October 2023
onward are peace movements. Before that date the category is not used — earlier
anti-occupation / joint Arab-Jewish solidarity protests were hand-labelled and must be
left as they are.

HAREDI JEWS is for Haredi collective action only (draft, autopsies, light rail,
Sabbath, "unkosher" stores). Chabad / messianic / religious-Zionist protests are right
wing protesters.

NOT RELEVANT also covers: labor disputes and municipal-merger protests by workers of a
settlement council; local residents protesting crime, infrastructure or services with
no political demand.

ARMY ACTS (hand review of the Fortress log, 2026-10-03):
An act carried out by soldiers, Border Police, a regional-defence (גמ״ר / hagmar) soldier, the Civil
Administration, a regional-council inspector or a settlement guard / security officer acting with
soldiers is `not relevant` — even when settlers are present, escort them, or asked for it
("soldiers accompanied by settlers dismantled a tent", "security forces ordered the family to leave
after settlers complained", "masked soldiers expelled a shepherd", "a hagmar soldier shot ...").
A settlement security coordinator (רבש״ץ) who himself shoots or attacks, without soldiers, is `settlers`.
A fire or attack whose perpetrator is not named at all ("a fire broke out in a house") is `not relevant`.
Unattributed killings during clashes in a village that settlers entered stay `settlers`.

INPUT LANGUAGE: the description may be a Hebrew line from a settler-violence log instead of ACLED
English; apply the same rules. Such lines are settler events unless they fit ARMY ACTS above or
name Jewish right-wing activists outside the West Bank (aid-convoy blockades, Flag March, Old City,
al-Aqsa) — those are `right wing protesters`.
"""

SCHEMA = {'type': 'object', 'properties': {'main_actor': {'type': 'string', 'enum': ['settlers', 'protesters against government', 'arab israelis', 'haredi jews', 'right wing protesters', 'peace movements', 'not relevant']}, 'certainty': {'type': 'string', 'enum': ['high', 'medium', 'low']}, 'needs_review': {'type': 'string', 'enum': ['yes', 'no']}, 'reason': {'type': 'string'}}, 'required': ['main_actor', 'certainty', 'needs_review', 'reason'], 'additionalProperties': False}



def norm(v):
    return "" if v is None else str(v).strip()

def hnorm(v):
    return norm(v).lower().replace(" ", "_")

def find_col(headers, names, required=False):
    lookup = {hnorm(v): i + 1 for i, v in enumerate(headers) if v is not None}
    for name in names:
        c = lookup.get(hnorm(name))
        if c:
            return c
    if required:
        raise KeyError(f"Missing required column. Tried: {names}")
    return None

def find_or_add_col(ws, name):
    headers = [c.value for c in ws[1]]
    c = find_col(headers, [name], required=False)
    if c:
        return c
    c = ws.max_column + 1
    ws.cell(1, c, name)
    return c

def extract_output_text(body):
    for item in body.get("output", []):
        if item.get("type") == "message":
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    return content.get("text", "")
    return ""

def save_state(base, state):
    (base / STATE_FILE).write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

def load_state(base):
    p = base / STATE_FILE
    if not p.exists():
        raise FileNotFoundError(f"{STATE_FILE} not found. Run submit first.")
    return json.loads(p.read_text(encoding="utf-8"))

def request_body(user_text):
    return {
        "model": MODEL,
        "reasoning": {"effort": REASONING_EFFORT},
        "instructions": INSTRUCTIONS,
        "input": user_text,
        "text": {
            "format": {
                "type": "json_schema",
                "name": "main_actor_result",
                "strict": True,
                "schema": SCHEMA,
            }
        },
        "max_output_tokens": 700,
    }

def submit_chunk(client, base, chunk):
    p = Path(chunk["input_path"])
    with p.open("rb") as f:
        uploaded = client.files.create(file=f, purpose="batch")
    batch = client.batches.create(
        input_file_id=uploaded.id,
        endpoint="/v1/responses",
        completion_window="24h",
    )
    chunk["input_file_id"] = uploaded.id
    chunk["batch_id"] = batch.id
    chunk["status"] = batch.status
    chunk["attempts"] = chunk.get("attempts", 0) + 1
    print(f"[{chunk['chunk_number']:03d}] submitted {chunk['request_count']:,} -> {batch.id} ({batch.status})")

def refresh(client, state):
    for chunk in state["chunks"]:
        if not chunk.get("batch_id"):
            chunk["status"] = "pending"
            continue
        b = client.batches.retrieve(chunk["batch_id"])
        chunk["status"] = b.status
        counts = getattr(b, "request_counts", None)
        if counts:
            chunk["complete"] = getattr(counts, "completed", 0) or 0
            chunk["failed"] = getattr(counts, "failed", 0) or 0
            chunk["total"] = getattr(counts, "total", 0) or 0
        if getattr(b, "output_file_id", None):
            chunk["output_file_id"] = b.output_file_id
        if getattr(b, "error_file_id", None):
            chunk["error_file_id"] = b.error_file_id
        if b.status == "failed":
            chunk["batch_errors"] = str(getattr(b, "errors", "") or "")

def active_count(state):
    return sum(
        1 for c in state["chunks"]
        if c.get("status") in {"validating", "in_progress", "finalizing", "cancelling"}
    )

def submit_pending(client, base, state):
    slots = max(0, MAX_ACTIVE_BATCHES - active_count(state))
    n = 0
    for chunk in state["chunks"]:
        if n >= slots:
            break
        if chunk.get("status", "pending") != "pending":
            continue
        submit_chunk(client, base, chunk)
        n += 1
    return n

def build_jobs(ws):
    headers = [c.value for c in ws[1]]
    row_id_col = find_col(headers, ["row_id", "row id"], required=False)
    desc_col = find_col(headers, ["Description", "description"], required=True)
    he_col = find_col(headers, ["description_he_medium"], required=False)
    actor_col = find_col(headers, ["main_actor"], required=False)
    hand_col = find_col(headers, ["main_actor_hand"], required=False)
    hidden_col = find_col(headers, ["hidden"], required=False)

    jobs = []
    for excel_row in range(2, ws.max_row + 1):
        g = lambda c: norm(ws.cell(excel_row, c).value) if c else ""
        row_id = g(row_id_col)
        # ROW SELECTION (loop-safe): only rows with no main_actor yet, never a
        # hand-set row (date in main_actor_hand), never a hidden row. Rows from a
        # Hebrew log have no English Description — their Hebrew line is sent.
        if g(actor_col) or g(hand_col) or g(hidden_col):
            continue
        description = g(desc_col) or g(he_col)
        if not description:
            continue
        cid = row_id or f"excel-row-{excel_row}"
        jobs.append({
            "custom_id": cid,
            "excel_row": excel_row,
            "row_id": row_id,
            "user_text": f"EVENT DESCRIPTION:\n{description}",
        })

    return jobs

def command_submit(filename):
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    source = Path(filename).expanduser().resolve()
    base = source.parent
    batch_dir = base / BATCH_DIR
    batch_dir.mkdir(exist_ok=True)

    wb = load_workbook(source)  # full load: ws.cell() is O(1); read_only made it O(rows) per call
    try:
        ws = wb.active
        jobs = build_jobs(ws)
    finally:
        wb.close()

    chunks = []
    for start in range(0, len(jobs), CHUNK_SIZE):
        part = jobs[start:start + CHUNK_SIZE]
        number = len(chunks) + 1
        jsonl = batch_dir / f"main_actor_{number:03d}.jsonl"
        with jsonl.open("w", encoding="utf-8") as f:
            for job in part:
                req = {
                    "custom_id": job["custom_id"],
                    "method": "POST",
                    "url": "/v1/responses",
                    "body": request_body(job["user_text"]),
                }
                f.write(json.dumps(req, ensure_ascii=False) + "\n")
        chunks.append({
            "chunk_number": number,
            "input_path": str(jsonl),
            "request_count": len(part),
            "custom_ids": [j["custom_id"] for j in part],
            "status": "pending",
            "batch_id": None,
            "attempts": 0,
        })

    state = {
        "source_workbook": str(source),
        "rows_submitted": len(jobs),
        "jobs": {j["custom_id"]: j for j in jobs},
        "chunks": chunks,
    }
    save_state(base, state)

    client = OpenAI()
    submit_pending(client, base, state)
    save_state(base, state)

    print(f"Prepared {len(jobs):,} rows in {len(chunks)} chunks.")
    print("Run periodically:")
    print(f"python3 {Path(__file__).name} status")

def command_status():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    newly = submit_pending(client, base, state)
    save_state(base, state)

    print(f"\nChecking {len(state['chunks'])} chunks...\n")
    for c in state["chunks"]:
        total = c.get("total", c["request_count"])
        print(
            f"[{c['chunk_number']:03d}] {c.get('status','pending'):<12} "
            f"{c.get('complete',0):,}/{total:,} complete, "
            f"{c.get('failed',0):,} failed"
        )
        if c.get("status") == "failed" and c.get("batch_errors"):
            print("   ", c["batch_errors"][:500])

    completed = sum(c.get("status") == "completed" for c in state["chunks"])
    failed = sum(c.get("status") == "failed" for c in state["chunks"])
    pending = sum(c.get("status") == "pending" for c in state["chunks"])

    print(f"\nCompleted chunks: {completed}/{len(state['chunks'])}")
    print(f"Pending chunks: {pending}")
    print(f"Failed chunks: {failed}")
    if newly:
        print(f"Submitted {newly} new chunk(s).")
    if failed:
        print(f"Run: python3 {Path(__file__).name} retry_failed")
    if completed == len(state["chunks"]):
        print(f"Run: python3 {Path(__file__).name} download")

def command_retry_failed():
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    reset = 0
    for c in state["chunks"]:
        if c.get("status") == "failed":
            c["batch_id"] = None
            c["status"] = "pending"
            c["input_file_id"] = None
            c["output_file_id"] = None
            c["error_file_id"] = None
            reset += 1
    submit_pending(client, base, state)
    save_state(base, state)
    print(f"Reset {reset} failed chunk(s).")

def parse_output(content):
    results, errors = {}, {}
    for line in content.decode("utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        cid = obj.get("custom_id")
        if obj.get("error"):
            errors[cid] = str(obj["error"])
            continue
        response = obj.get("response") or {}
        if response.get("status_code") != 200:
            errors[cid] = f"HTTP {response.get('status_code')}"
            continue
        text = extract_output_text(response.get("body") or {})
        if not text:
            errors[cid] = "No output_text"
            continue
        try:
            results[cid] = json.loads(text)
        except Exception as exc:
            errors[cid] = f"JSON parse error: {exc}"
    return results, errors

def command_download():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    save_state(base, state)

    incomplete = [c for c in state["chunks"] if c.get("status") != "completed"]
    if incomplete:
        raise RuntimeError("Not all chunks are completed.")

    results, errors = {}, {}
    for i, c in enumerate(state["chunks"], 1):
        b = client.batches.retrieve(c["batch_id"])
        print(f"[{i:03d}/{len(state['chunks']):03d}] downloading...")
        content = client.files.content(b.output_file_id)
        parsed, errs = parse_output(content.content)
        results.update(parsed)
        errors.update(errs)

    source = Path(state["source_workbook"])
    wb = load_workbook(source)
    try:
        ws = wb.active
        actor_col = find_or_add_col(ws, "main_actor")
        certainty_col = find_or_add_col(ws, "main_actor_certainty")
        review_col = find_or_add_col(ws, "main_actor_needs_review")
        reason_col = find_or_add_col(ws, "main_actor_reason")

        for cid, job in state["jobs"].items():
            row = job["excel_row"]
            result = results.get(cid)
            if not result:
                continue
            ws.cell(row, actor_col, result["main_actor"])
            ws.cell(row, certainty_col, result["certainty"])
            ws.cell(row, review_col, result["needs_review"])
            ws.cell(row, reason_col, result["reason"])

        output = source.parent / (source.stem + " - actors" + source.suffix)
        wb.save(output)
    finally:
        wb.close()

    print(f"\nFinished. Valid results: {len(results):,}")
    print(f"Errors: {len(errors):,}")
    print(f"Output:\n{output}")

def main():
    if len(sys.argv) < 2:
        raise SystemExit(
            f'Usage:\n'
            f'  python3 {Path(__file__).name} submit "table.xlsx"\n'
            f'  python3 {Path(__file__).name} status\n'
            f'  python3 {Path(__file__).name} retry_failed\n'
            f'  python3 {Path(__file__).name} download'
        )
    cmd = sys.argv[1].lower()
    if cmd == "submit":
        if len(sys.argv) != 3:
            raise SystemExit("submit requires an .xlsx filename")
        command_submit(sys.argv[2])
    elif cmd == "status":
        command_status()
    elif cmd == "retry_failed":
        command_retry_failed()
    elif cmd == "download":
        command_download()
    else:
        raise SystemExit(f"Unknown command: {cmd}")

if __name__ == "__main__":
    main()
