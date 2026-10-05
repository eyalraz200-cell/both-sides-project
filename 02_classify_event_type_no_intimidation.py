#!/usr/bin/env python3
# Both Sides project
# Requirements:
#   pip install openai openpyxl
#   export OPENAI_API_KEY="..."

import json
import os
import sys
from pathlib import Path
from openai import OpenAI
from openpyxl import load_workbook

MODEL = "gpt-5.6-terra"
REASONING_EFFORT = "medium"
CHUNK_SIZE = 250
MAX_ACTIVE_BATCHES = 8

STATE_FILE = "event_type_state.json"
BATCH_DIR = "event_type_batches"

INSTRUCTIONS = r"""
You are classifying the TYPE OF EVENT.

You receive exactly three pieces of information:

1. MAIN ACTOR — the actor already identified as responsible for the event.
2. EVENT DESCRIPTION — the original description of what happened (English, or Hebrew for rows
   imported from a Hebrew source log; classify Hebrew text by the same rules).
3. CURRENT EVENT TYPE — may be blank and is supplied only as context.

The MAIN ACTOR value is authoritative.

Do NOT:
- infer a different main actor
- rename the main actor
- substitute another actor category
- classify according to a secondary actor

Your task is to classify ONLY the actions performed by the supplied MAIN ACTOR.


CRITICAL MAIN-ACTOR RULE

A description may mention several actors performing different actions.

Ignore actions performed by secondary actors when choosing the event type.

Even if a secondary actor performs a more violent or more severe action, that action must NOT determine the classification.

Examples:

- MAIN ACTOR = protesters against government.
  Protesters block a road. Police later use force to disperse them.
  -> חסימת כביש

- MAIN ACTOR = settlers.
  Settlers throw stones at vehicles. Soldiers later fire shots.
  -> תקיפה בנשק קר

- MAIN ACTOR = haredi jews.
  Haredi protesters demonstrate. Police later use force.
  -> הפגנה לא אלימה, unless the Haredi protesters themselves performed a more severe action.

Severity is applied ONLY among actions performed by the MAIN ACTOR.


ALLOWED EVENT TYPES

1. הפגנה לא אלימה
2. חסימת כביש
3. הפרות סדר
4. פגיעה ברכוש
5. ניכוס שטח
6. תקיפה פיזית
7. החזקה בכפייה
8. תקיפה בנשק קר
9. תקיפה בנשק חם
10. פוגרום


1. הפגנה לא אלימה

A demonstration, rally, march, vigil, protest gathering, or protest presence in which the MAIN ACTOR does not perform any more severe action.

Use this only when the MAIN ACTOR itself does not:
- block roads
- riot or engage in disorder
- damage or steal property
- appropriate land
- physically attack people
- forcibly hold people
- use a weapon

Speech by itself does NOT escalate the event.

The following, by themselves, may remain הפגנה לא אלימה:
- slogans
- insults
- hostile chants
- racist chants
- calling someone a traitor
- inflammatory political speech
- generalized chants such as "death to..." when they are slogans rather than a direct targeted threat or attack

Statements alone are not a separate event type.


2. חסימת כביש

The MAIN ACTOR blocks a:
- road
- highway
- intersection
- junction
- entrance
- access road
- traffic route

as part of a protest or confrontation.

A symbolic protest camp, tent, structure, or "outpost" whose purpose is protest or road obstruction is NOT ניכוס שטח.

If its purpose is to obstruct traffic or access, classify it as חסימת כביש.

If intimidation or threats occur during an event whose central concrete action is blocking traffic/access, the road-blocking action remains חסימת כביש unless the MAIN ACTOR also performs a more severe concrete action.


3. הפרות סדר

The MAIN ACTOR participates in:
- riots
- clashes
- violent disorder
- disruptive confrontations
- public disorder

Use this when the conduct goes beyond a peaceful demonstration but does not meet a more severe category.

Generic wording such as:
- "clashes occurred"
- "fighting broke out"
- "people were injured"
- "objects were thrown"

does NOT by itself prove a physical or armed attack by the MAIN ACTOR.

If the description does not explicitly establish a more severe action by the MAIN ACTOR, use הפרות סדר.

Throwing or igniting:
- fireworks
- firecrackers
- stun grenades
- bottles
- tires
- bins
- miscellaneous objects

without a clearly identified human target should normally be classified as הפרות סדר, unless another rule below applies.

Burning tires, bins, or similar objects as a protest tactic is not automatically פגיעה ברכוש.
If done to block a road -> חסימת כביש.
If done as part of disorder -> הפרות סדר.


THREATS / INTIMIDATION RULE — THERE IS NO SEPARATE INTIMIDATION CATEGORY

Do NOT output "הטרדה ואיומים". It is not an allowed category.

Classify intimidation according to what the MAIN ACTOR actually did:

A. Speech-only protest rhetoric
Insults, hostile slogans, racist chants, political abuse, or generalized threatening rhetoric during an otherwise peaceful demonstration:
-> הפגנה לא אלימה

B. Direct targeted intimidation without a more concrete action
Explicit threats directed at an identifiable person/community, coercive intimidation, following/stalking, chasing someone in order to intimidate them, sustained menacing conduct, or threatening presence that goes beyond ordinary protest speech:
-> הפרות סדר

Examples:
- directly threatening residents that they will be harmed if they do not leave -> הפרות סדר
- repeatedly following and intimidating a journalist -> הפרות סדר
- surrounding a person in a menacing confrontation without touching them -> הפרות סדר

C. Intimidation through a concrete action
If the intimidation is carried out through another category, classify the concrete action:
- blocking a road/access -> חסימת כביש
- damaging property -> פגיעה ברכוש
- taking over land -> ניכוס שטח
- hitting/pushing/spitting or other direct bodily attack -> תקיפה פיזית
- forcibly holding someone -> החזקה בכפייה
- threatening/attacking a person with a cold weapon -> תקיפה בנשק קר
- pointing/using a firearm, explosive, Molotov cocktail, or other hot weapon directly at people -> תקיפה בנשק חם

The existence of fear, intimidation, or threats does not override a more specific concrete action.


4. פגיעה ברכוש

The MAIN ACTOR intentionally damages, destroys, steals, loots, or takes property.

This includes:
- homes
- vehicles
- businesses
- equipment
- infrastructure
- agricultural property
- crops
- privately owned objects
- theft
- looting
- taking property

IMPORTANT EXCLUSIONS:

Do NOT classify an action as פגיעה ברכוש merely because the MAIN ACTOR:
- removes or destroys a protest sign
- tears down a banner
- moves or knocks over a temporary barrier
- removes a barricade
- dismantles a temporary roadblock
- moves a checkpoint barrier
- damages another temporary public-order object

When these actions are part of a protest or confrontation, classify the broader conduct instead, usually:
- הפרות סדר
or
- חסימת כביש

STONE RULE:
Throwing stones at a person, at a vehicle, OR AT A HOUSE / BUILDING / SHOP is NOT פגיעה ברכוש.
It is תקיפה בנשק קר — even when nobody is reported hit and even outside the West Bank
(stones at an apartment, an office window, a home: still תקיפה בנשק קר).

ARSON RULE:
Deliberately burning an EMPTY building, EMPTY vehicle, field, crop, or other property, with no people directly endangered, is פגיעה ברכוש.

THEFT RULE:
Theft, looting, or taking property may be classified as פגיעה ברכוש even when the item is not physically damaged.

LIVESTOCK RULE:
Killing, injuring, stealing, or running over animals is פגיעה ברכוש.

INJURY RULE:
If the MAIN ACTOR's property destruction injured a person (someone inside the house that
was destroyed, someone hurt "during the attack"), the event is at least תקיפה פיזית.


5. ניכוס שטח

Use this ONLY when the MAIN ACTOR is actually establishing or expanding territorial control over land that does not belong to the acting group.

Examples:
- establishing a permanent or territorial outpost
- fencing off land
- taking over land
- cultivating land in order to establish control
- constructing or placing structures to establish territorial possession
- expanding control over a plot or area

There must be a genuine territorial-control purpose.

Do NOT use ניכוס שטח for:
- GRAZING INCURSIONS: settlers bringing herds onto Palestinian land, grazing sheep /
  cattle / camels in fields or between houses, however often it recurs -> הפרות סדר
  (or פגיעה ברכוש only if crops are explicitly destroyed, תקיפה פיזית if people are
  attacked). ניכוס שטח needs a tent, caravan, fence, outpost, ploughing or structures.
- a protest encampment
- a symbolic "democracy outpost"
- tents erected temporarily for a demonstration
- occupying a road as a protest tactic
- blocking access
- temporary protest structures

Those should be classified according to the protest action, usually חסימת כביש, הפרות סדר, or הפגנה לא אלימה.


6. תקיפה פיזית

The MAIN ACTOR explicitly performs direct bodily violence against a person WITHOUT a weapon.

Examples:
- punches
- kicks
- pushes
- shoves
- strikes with bare hands
- spits directly on a person
- physically beats someone without a weapon

The physical attack must be explicitly attributed to the MAIN ACTOR.

Do NOT infer תקיפה פיזית merely because:
- "a fight broke out"
- "clashes occurred"
- "a confrontation occurred"
- someone was injured
- violence occurred without identifying what the MAIN ACTOR did

In those ambiguous cases, prefer הפרות סדר.

HAND-REVIEW RULES (Fortress log, 2026-10-03):
- Taking over an inhabited house while the family is inside -> החזקה בכפייה.
- Arson of houses, structures, vehicles or between houses (fire SET, not gunfire) -> פגיעה ברכוש,
  never תקיפה בנשק חם (that is for live fire only).
- "מטען" on an aid truck / convoy means CARGO, not an explosive: dumping or throwing the cargo is
  פגיעה ברכוש.
- Waste, rubbish or "objects" thrown at a house -> הפרות סדר.
- A flag raised on a hill, with nothing built -> הפרות סדר (not ניכוס שטח).
- A record where settlers entered and people were killed or wounded by gunfire, and the SOURCE TAGS
  include "ירי לעבר אזרחים" or "הריגה" -> תקיפה בנשק חם, even if the text does not name the shooter.
  Without such tags, unattributed gunfire during "clashes" or "an exchange of fire" stays הפרות סדר.
- A herd brought near houses or into a yard, a minor leading goats from an outpost -> הפרות סדר.
- Protesters who BLOCK a road AND clash with police, light bonfires or set up barricades
  -> הפרות סדר (not חסימת כביש), even when the police also used force. Blocking alone, or
  blocking with arrests only, stays חסימת כביש.
- Protesters who CLASHED WITH POLICE ("clashed with police", "scuffled with officers",
  "confrontations with police") -> הפרות סדר, with or without a road block. Arrests, detentions or
  police dispersal ALONE, with no clash described, do not raise the category.
- Levelling / bulldozing Palestinian land to pave or open a road for a settlement or outpost
  -> ניכוס שטח (the land is taken), even when a water line or crops are destroyed on the way.
  Bulldozing, levelling or uprooting land with NO road, outpost, fence or structure stated
  -> פגיעה ברכוש (it destroys the land, it does not take it).
- Molotov cocktails thrown at INHABITED houses, and setting fire to an OCCUPIED vehicle (people
  inside), are תקיפה בנשק חם — fire aimed at people. Arson of empty property stays פגיעה ברכוש.
- עימותים / "confrontations" / "scuffles" / "heated arguments that turned physical" / "minor
  shoving" are NOT תקיפה פיזית. תקיפה פיזית needs a DESCRIBED assault by the main actor on a
  named target ("beat a man", "pushed the officer to the ground", "kicked him"). A protest
  with confrontations and nothing more described is הפרות סדר — whether or not anyone was
  arrested, and whatever the arrest was "on suspicion of".
- When the PROTESTERS were the ones attacked (rioters / passers-by / counter-protesters hit
  them), classify the protesters' own conduct only (usually הפגנה לא אלימה).
- One participant's act inside a mass demonstration ("a protester slapped an MK") does not
  make the crowd's event an assault -> הפרות סדר at most.
- Protesters who "clashed with counter-protesters", with no violent act described for the
  MAIN ACTOR themselves -> הפגנה לא אלימה. Only an act the main actor did (hit, threw, pushed)
  raises it.

PASSIVE-VOICE RULE (source logs written from the victims' side):
"A Palestinian was attacked / shot / run over", "stones were thrown", "a house was set on
fire" with no attacker named means the MAIN ACTOR did it — classify the act, do not
drop to a lower category because the sentence has no subject. A settlement security
guard, security coordinator (רבש"ץ) or "settlers in uniform" count as the main actor.

DISPLACEMENT-OUTCOME RULE:
"The family left / the community was displaced after repeated attacks" with no concrete
act described in THIS record is הפרות סדר. Forcing someone out of their house and
settlers moving in is ניכוס שטח.

"ATTACKED RESIDENTS" RULE:
"Attacked residents / civilians / farmers / shepherds" with no mechanism given IS
תקיפה פיזית (people were the target). "Attacked the village / the community / houses"
with no mechanism and no people mentioned is הפרות סדר.

INCURSION RULE:
Settlers who enter a Palestinian village (with or without army escort) and clashes
follow are הפרות סדר, never הפגנה לא אלימה — even when the description only says the
Palestinians rioted.

SYMBOLIC CROSSING RULE:
Crossing into Gaza / breaking through a crossing to plant trees or "start a settlement"
as a protest stunt is הפרות סדר, not ניכוס שטח.


7. החזקה בכפייה

The MAIN ACTOR successfully:
- abducts
- kidnaps
- captures
- forcibly detains
- holds a person against their will

The person must actually be taken or held.

An ATTEMPTED kidnapping or attempted abduction that fails is NOT החזקה בכפייה.

Seizing someone briefly and marching or handing them over to soldiers / police is NOT
החזקה בכפייה — classify the assault that came with it (usually תקיפה פיזית). Use
החזקה בכפייה only when the person was held for a sustained period (hours or more) or
taken away to another location and kept there.
A completed ABDUCTION / KIDNAPPING ("kidnapped two young men, assaulted them and released
them") is החזקה בכפייה even when the victims were released soon after. Detaining people IN
PLACE — stopping, binding hands, blindfolding, questioning, beating — and then letting them go,
with no taking away, is תקיפה פיזית (or the weapon category if one was used).

Classify an unsuccessful attempt according to what the MAIN ACTOR actually did.


8. תקיפה בנשק קר

The MAIN ACTOR attacks or directly threatens a person using a cold weapon, dangerous hand-held object, or specified lower-level projectile.

This includes:
- knife
- blade
- club
- stick
- metal bar
- blunt object
- sharp object
- stones
- rocks

STONE RULE:
Throwing stones at PEOPLE, VEHICLES, OR HOUSES / BUILDINGS is always תקיפה בנשק קר.
If the main actor's stones injured anyone, it is at least תקיפה בנשק קר regardless of
what else the description calls property damage.

Also classify the following as תקיפה בנשק קר ONLY when they are clearly directed at people:
- thrown bottles
- stun grenades
- pepper spray, taser (an irritant or shock device used on a person)
(NOT fireworks, firecrackers or burning objects aimed at people — those are תקיפה בנשק חם,
see the FIRE AT PEOPLE rule below.)
- a firearm used as a club (rifle butt, pistol whip) — not fired, so a blunt object

NOT תקיפה בנשק קר, even when thrown at people (classify as הפרות סדר unless something
more severe applies):
- eggs
- smoke grenades
- "objects" / "various objects" with nothing more specific

Examples:
- bottle thrown at a person -> תקיפה בנשק קר
- stun grenade thrown toward people -> תקיפה בנשק קר
- fireworks fired directly at people -> תקיפה בנשק חם (hand rule 2026-10-04)
- firecrackers thrown at people -> תקיפה בנשק חם
- a burning torch / flare / Molotov cocktail thrown at a person -> תקיפה בנשק חם
- knife brandished directly at a person as a threat -> תקיפה בנשק קר

But:
- fireworks simply set off during disorder -> הפרות סדר
- firecrackers thrown with no identified human target -> הפרות סדר

FIRE AT PEOPLE RULE (hand review 2026-10-04):
Anything that burns or explodes, LAUNCHED OR THROWN AT A PERSON — a firework, a firecracker,
a flare, a burning torch, a Molotov cocktail — is תקיפה בנשק חם, the same class as gunfire,
whether or not it hit. Fireworks fired DURING A RIOT OR CLASH WITH POLICE ("threw stones and
fireworks and clashed with police", "threw stones at officers, fired fireworks") count as aimed
at the police -> תקיפה בנשק חם. Only fireworks set off with no confrontation around them
("fireworks were set off during the celebration / while chasing people out") stay as they are.
- stun grenade used with no identified human target -> הפרות סדר
- bottles thrown generically during clashes with no clear human target -> הפרות סדר

UNSPECIFIED OBJECT RULE:
"Threw objects" or "objects were thrown" is NOT enough for תקיפה בנשק קר unless the object and relevant target are identified.

RAMMING RULE:
Any vehicle (car, truck, motorcycle, ATV, tractor, bulldozer) deliberately driven at a
person OR at an occupied vehicle is תקיפה בנשק קר. An ATTEMPT with no contact is still
תקיפה בנשק קר. Running over livestock is פגיעה ברכוש (animals are property).
A passing driver who hits demonstrators is a secondary actor — it never classifies the
demonstrators' event.


9. תקיפה בנשק חם

The MAIN ACTOR attacks or directly threatens people using:

- firearm
- handgun
- rifle
- gunfire
- live ammunition
- fragmentation grenade
- explosive grenade
- explosive device
- bomb
- serious explosive material
- Molotov cocktail / firebomb

GRENADE RULE:
- fragmentation grenade / explosive grenade / ordinary lethal grenade -> תקיפה בנשק חם
- stun grenade -> תקיפה בנשק קר only when directed at people; otherwise usually הפרות סדר

FIREARM-THREAT RULE:
Pointing or aiming a firearm directly at a person as a threat is תקיפה בנשק חם even if no shot is fired.

"ARMED" RULE:
Being described as "armed" is NOT enough. If the armed actors did not fire, point, or
strike with the weapon, classify the act they actually performed ("armed settlers attacked
civilians, injuring 2" -> תקיפה פיזית).

FIRE-AT-HOUSES RULE:
Live fire toward inhabited houses, tents, or a village is תקיפה בנשק חם even when no
individual is named as the target.

CHEMICAL RULE:
Acid, chlorine, an unidentified gas or liquid sprayed or poured on people or into an
occupied dwelling, and TEAR GAS CANISTERS FIRED from a launcher are תקיפה בנשק חם.
Hand-held pepper spray and tasers stay תקיפה בנשק קר.

ARSON RULE:
Ordinary deliberate arson is תקיפה בנשק חם ONLY when:
- people are inside the targeted building or vehicle
OR
- the fire is directly used to attack or endanger people

If an empty building, empty vehicle, field, crop, or property is burned without people directly endangered -> פגיעה ברכוש.


10. פוגרום

Use פוגרום only when the MAIN ACTOR is explicitly a LARGE CROWD, MOB, MASS GATHERING, or similarly large collective that attacks:
- a community
- neighborhood
- village
- residential area
- civilian population

and the attack includes violence against people, homes, vehicles, businesses, or the civilian environment.

פוגרום requires ALL FOUR, performed by the MAIN ACTOR and EACH STATED EXPLICITLY in the text
(hand rule 2026-10-04 — a pogrom is a very heavy accusation; if any element has to be
inferred, it is not a pogrom):
1. A CROWD OF 30 OR MORE — "dozens", "hundreds", "thousands", or a stated number of 30+.
   (15 or 20 attackers, "tens", "a group", "several settlers", or NO size stated at all is
   NOT a crowd for this purpose — hand rule 2026-10-04.)
2. VIOLENCE AGAINST PEOPLE — beatings, stone-throwing at people or inhabited houses,
   gunfire, pepper spray — with victims present.
3. FIRE — houses, tents, vehicles, fields or buildings in the community SET ON FIRE.
4. PROPERTY DESTROYED — the fire or other acts damaged or destroyed homes, vehicles,
   businesses or crops (stated, not assumed).

All four during a raid on a village, community, neighbourhood or farm = פוגרום
(e.g. 50 settlers beat residents injuring 2 AND burn 7 vehicles -> פוגרום; the same by
20 settlers -> תקיפה פיזית, by a crowd of unstated size -> classify the acts themselves).

NOT a pogrom:
- a crowd that only burns EMPTY property with nobody attacked -> פגיעה ברכוש
- a crowd that only beats or stones people with no fire -> תקיפה פיזית / תקיפה בנשק קר
- a mass march or demonstration in a city that turns violent (a flag march where
  marchers beat passers-by) -> classify the violence itself
- "attacked a community" with no mechanism -> הפרות סדר
- a small group (under 30) however violent, or a raid whose size is not stated

When all four conditions hold, פוגרום is the answer (it is top of the hierarchy).

POGROM STANDARD (user, 2026-10-04 — a pogrom is a very heavy accusation; when in doubt, it is
NOT a pogrom). ALL of these must be stated in the text itself:
- a crowd of 30+ with its size given (no stated size = not a pogrom, however severe);
- Palestinians injured or killed BY THE MAIN ACTOR (injuries caused only by soldiers/police, or
  only settlers hurt, do not count; "no casualties" rules it out);
- fire set to homes, vehicles, structures or fields — named, not just "set fire to property";
- enough concrete detail to judge (a one-line "wave of attacks including arson, stones and
  assaults" with no specifics is NOT enough).
Missing any one -> classify the most serious act described (gunfire -> תקיפה בנשק חם, stones or
clubs at people -> תקיפה בנשק קר, beatings -> תקיפה פיזית, arson of property -> פגיעה ברכוש).


SEVERITY HIERARCHY

If the MAIN ACTOR performs more than one applicable action, classify the event according to the highest-ranked action performed BY THAT MAIN ACTOR:

1 — הפגנה לא אלימה
2 — חסימת כביש
3 — הפרות סדר
4 — פגיעה ברכוש
5 — ניכוס שטח
6 — תקיפה פיזית
7 — החזקה בכפייה
8 — תקיפה בנשק קר
9 — תקיפה בנשק חם
10 — פוגרום


FINAL CHECK BEFORE ANSWERING

Before returning the category, verify:

1. Am I classifying an action performed by the supplied MAIN ACTOR?
2. Did I accidentally use an action performed by a secondary actor?
3. If several MAIN ACTOR actions occurred, did I choose the highest-ranked applicable category?
4. Did I apply the special rules for threats/intimidation, stones (incl. at houses), bottles, stun grenades, fireworks, firecrackers, eggs/smoke/unspecified objects, pepper spray vs. chemicals vs. tear-gas launchers, "armed" without use, fire at houses, arson, injuries from property destruction, livestock, protest barriers, land appropriation vs. symbolic crossings, incursions into villages, brief seizure vs. sustained holding, ramming, and the four-part pogrom test?

Return exactly one allowed event type.

"""

SCHEMA = {'type': 'object', 'properties': {'event_type': {'type': 'string', 'enum': ['הפגנה לא אלימה', 'חסימת כביש', 'הפרות סדר', 'פגיעה ברכוש', 'ניכוס שטח', 'תקיפה פיזית', 'החזקה בכפייה', 'תקיפה בנשק קר', 'תקיפה בנשק חם', 'פוגרום']}, 'certainty': {'type': 'string', 'enum': ['high', 'medium', 'low']}, 'needs_review': {'type': 'string', 'enum': ['yes', 'no']}, 'reason': {'type': 'string'}}, 'required': ['event_type', 'certainty', 'needs_review', 'reason'], 'additionalProperties': False}



def xl_safe(v):
    """Strip control characters Excel refuses (openpyxl raises IllegalCharacterError)."""
    if not isinstance(v, str):
        return v
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", v)

def norm(v):
    return "" if v is None else str(v).strip()

def hnorm(v):
    return norm(v).lower().replace(" ", "_")

def find_col(headers, names, required=False):
    lookup = {hnorm(v): i + 1 for i, v in enumerate(headers) if v is not None}
    for name in names:
        c = lookup.get(hnorm(name))
        if c:
            return c
    if required:
        raise KeyError(f"Missing required column. Tried: {names}")
    return None

def find_or_add_col(ws, name):
    headers = [c.value for c in ws[1]]
    c = find_col(headers, [name], required=False)
    if c:
        return c
    c = ws.max_column + 1
    ws.cell(1, c, name)
    return c

def extract_output_text(body):
    for item in body.get("output", []):
        if item.get("type") == "message":
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    return content.get("text", "")
    return ""

def save_state(base, state):
    (base / STATE_FILE).write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

def load_state(base):
    p = base / STATE_FILE
    if not p.exists():
        raise FileNotFoundError(f"{STATE_FILE} not found. Run submit first.")
    return json.loads(p.read_text(encoding="utf-8"))

def request_body(user_text):
    return {
        "model": MODEL,
        "reasoning": {"effort": REASONING_EFFORT},
        "instructions": INSTRUCTIONS,
        "input": user_text,
        "text": {
            "format": {
                "type": "json_schema",
                "name": "event_type_result",
                "strict": True,
                "schema": SCHEMA,
            }
        },
        "max_output_tokens": 700,
    }

def submit_chunk(client, base, chunk):
    p = Path(chunk["input_path"])
    with p.open("rb") as f:
        uploaded = client.files.create(file=f, purpose="batch")
    batch = client.batches.create(
        input_file_id=uploaded.id,
        endpoint="/v1/responses",
        completion_window="24h",
    )
    chunk["input_file_id"] = uploaded.id
    chunk["batch_id"] = batch.id
    chunk["status"] = batch.status
    chunk["attempts"] = chunk.get("attempts", 0) + 1
    print(f"[{chunk['chunk_number']:03d}] submitted {chunk['request_count']:,} -> {batch.id} ({batch.status})")

def refresh(client, state):
    for chunk in state["chunks"]:
        if not chunk.get("batch_id"):
            chunk["status"] = "pending"
            continue
        b = client.batches.retrieve(chunk["batch_id"])
        chunk["status"] = b.status
        counts = getattr(b, "request_counts", None)
        if counts:
            chunk["complete"] = getattr(counts, "completed", 0) or 0
            chunk["failed"] = getattr(counts, "failed", 0) or 0
            chunk["total"] = getattr(counts, "total", 0) or 0
        if getattr(b, "output_file_id", None):
            chunk["output_file_id"] = b.output_file_id
        if getattr(b, "error_file_id", None):
            chunk["error_file_id"] = b.error_file_id
        if b.status == "failed":
            chunk["batch_errors"] = str(getattr(b, "errors", "") or "")

def active_count(state):
    return sum(
        1 for c in state["chunks"]
        if c.get("status") in {"validating", "in_progress", "finalizing", "cancelling"}
    )

def submit_pending(client, base, state):
    slots = max(0, MAX_ACTIVE_BATCHES - active_count(state))
    n = 0
    for chunk in state["chunks"]:
        if n >= slots:
            break
        if chunk.get("status", "pending") != "pending":
            continue
        submit_chunk(client, base, chunk)
        n += 1
    return n

def build_jobs(ws):
    _h = [c.value for c in ws[1]]
    _src = find_col(_h, ["source"], required=False)
    PLO_ONLY = {"plo negotiations affairs department"}     # server.py SOLE_SOURCE_EXCLUDE: never shipped
    _corr = find_col(_h, ["corroborated_by"], required=False)
    def _plo_only(row):
        if not _src: return False
        if _corr and norm(ws.cell(row, _corr).value): return False   # corroborated by the Fortress log: ships
        s = {x.strip().lower() for x in norm(ws.cell(row, _src).value).split(";") if x.strip()}
        return bool(s) and s <= PLO_ONLY
    headers = [c.value for c in ws[1]]
    row_id_col = find_col(headers, ["row_id", "row id"], required=False)
    actor_col = find_col(headers, ["main_actor", "main actor"], required=True)
    desc_col = find_col(headers, ["Description", "description"], required=True)
    current_col = find_col(headers, ["event_type", "event type"], required=False)
    hand_col = find_col(headers, ["event_type_hand"], required=False)
    recl_col = find_col(headers, ["reclassify"], required=False)
    he_col = find_col(headers, ["description_he_medium"], required=False)
    tags_col = find_col(headers, ["fortress_categories"], required=False)
    hidden_col = find_col(headers, ["hidden"], required=False)

    jobs = []
    for excel_row in range(2, ws.max_row + 1):
        if _plo_only(excel_row):
            continue
        row_id = norm(ws.cell(excel_row, row_id_col).value) if row_id_col else ""
        actor = norm(ws.cell(excel_row, actor_col).value)
        description = norm(ws.cell(excel_row, desc_col).value)
        current = norm(ws.cell(excel_row, current_col).value) if current_col else ""
        # Rows imported from a Hebrew log (the Fortress) have no English Description;
        # their text is in description_he_medium, and their own tags help as context.
        if not description and he_col:
            description = norm(ws.cell(excel_row, he_col).value)
        tags = norm(ws.cell(excel_row, tags_col).value) if tags_col else ""
        if not actor or not description or actor == "not relevant":
            continue
        if hidden_col and norm(ws.cell(excel_row, hidden_col).value):
            continue                      # hidden rows (ACLED duplicates etc.) never ship
        # ROW SELECTION: classify only rows that still need a type — blank event_type,
        # or reclassify = yes (set by 04 after a split). A row with a date in
        # event_type_hand is a hand decision and is never re-run unless reclassify says so.
        recl = norm(ws.cell(excel_row, recl_col).value).lower() == "yes" if recl_col else False
        hand = norm(ws.cell(excel_row, hand_col).value) if hand_col else ""
        # RERUN_ALL=1: the full second pass — every row that already has a type is
        # re-sent too, EXCEPT hand-set ones. download keeps the old value in
        # event_type_prev so the changes can be reviewed before shipping.
        if os.getenv("RERUN_ALL") == "1":
            if hand and not recl:
                continue
        elif not recl and (current or hand):
            continue
        cid = row_id or f"excel-row-{excel_row}"
        jobs.append({
            "custom_id": cid,
            "excel_row": excel_row,
            "row_id": row_id,
            "user_text": (
                f"MAIN ACTOR:\n{actor}\n\n"
                f"CURRENT EVENT TYPE:\n{current or '(blank)'}\n\n"
                + (f"SOURCE TAGS (the original log's own labels, context only):\n{tags}\n\n" if tags else "")
                + f"EVENT DESCRIPTION:\n{description}"
            ),
        })

    return jobs

def command_submit(filename):
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    source = Path(filename).expanduser().resolve()
    base = source.parent
    batch_dir = base / BATCH_DIR
    batch_dir.mkdir(exist_ok=True)

    wb = load_workbook(source)  # full load: ws.cell() is O(1); read_only made it O(rows) per call
    try:
        ws = wb.active
        jobs = build_jobs(ws)
    finally:
        wb.close()

    chunks = []
    for start in range(0, len(jobs), CHUNK_SIZE):
        part = jobs[start:start + CHUNK_SIZE]
        number = len(chunks) + 1
        jsonl = batch_dir / f"event_type_{number:03d}.jsonl"
        with jsonl.open("w", encoding="utf-8") as f:
            for job in part:
                req = {
                    "custom_id": job["custom_id"],
                    "method": "POST",
                    "url": "/v1/responses",
                    "body": request_body(job["user_text"]),
                }
                f.write(json.dumps(req, ensure_ascii=False) + "\n")
        chunks.append({
            "chunk_number": number,
            "input_path": str(jsonl),
            "request_count": len(part),
            "custom_ids": [j["custom_id"] for j in part],
            "status": "pending",
            "batch_id": None,
            "attempts": 0,
        })

    state = {
        "source_workbook": str(source),
        "rows_submitted": len(jobs),
        "jobs": {j["custom_id"]: j for j in jobs},
        "chunks": chunks,
    }
    save_state(base, state)

    client = OpenAI()
    submit_pending(client, base, state)
    save_state(base, state)

    print(f"Prepared {len(jobs):,} rows in {len(chunks)} chunks.")
    print("Run periodically:")
    print(f"python3 {Path(__file__).name} status")

def command_status():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    newly = submit_pending(client, base, state)
    save_state(base, state)

    print(f"\nChecking {len(state['chunks'])} chunks...\n")
    for c in state["chunks"]:
        total = c.get("total", c["request_count"])
        print(
            f"[{c['chunk_number']:03d}] {c.get('status','pending'):<12} "
            f"{c.get('complete',0):,}/{total:,} complete, "
            f"{c.get('failed',0):,} failed"
        )
        if c.get("status") == "failed" and c.get("batch_errors"):
            print("   ", c["batch_errors"][:500])

    completed = sum(c.get("status") == "completed" for c in state["chunks"])
    failed = sum(c.get("status") == "failed" for c in state["chunks"])
    pending = sum(c.get("status") == "pending" for c in state["chunks"])

    print(f"\nCompleted chunks: {completed}/{len(state['chunks'])}")
    print(f"Pending chunks: {pending}")
    print(f"Failed chunks: {failed}")
    if newly:
        print(f"Submitted {newly} new chunk(s).")
    if failed:
        print(f"Run: python3 {Path(__file__).name} retry_failed")
    if completed == len(state["chunks"]):
        print(f"Run: python3 {Path(__file__).name} download")

def command_retry_failed():
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    reset = 0
    for c in state["chunks"]:
        if c.get("status") == "failed":
            c["batch_id"] = None
            c["status"] = "pending"
            c["input_file_id"] = None
            c["output_file_id"] = None
            c["error_file_id"] = None
            reset += 1
    submit_pending(client, base, state)
    save_state(base, state)
    print(f"Reset {reset} failed chunk(s).")

def parse_output(content):
    results, errors = {}, {}
    for line in content.decode("utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        cid = obj.get("custom_id")
        if obj.get("error"):
            errors[cid] = str(obj["error"])
            continue
        response = obj.get("response") or {}
        if response.get("status_code") != 200:
            errors[cid] = f"HTTP {response.get('status_code')}"
            continue
        text = extract_output_text(response.get("body") or {})
        if not text:
            errors[cid] = "No output_text"
            continue
        try:
            results[cid] = json.loads(text)
        except Exception as exc:
            errors[cid] = f"JSON parse error: {exc}"
    return results, errors

def command_download():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    save_state(base, state)

    incomplete = [c for c in state["chunks"] if c.get("status") != "completed"]
    if incomplete:
        raise RuntimeError("Not all chunks are completed.")

    results, errors = {}, {}
    for i, c in enumerate(state["chunks"], 1):
        b = client.batches.retrieve(c["batch_id"])
        print(f"[{i:03d}/{len(state['chunks']):03d}] downloading...")
        content = client.files.content(b.output_file_id)
        parsed, errs = parse_output(content.content)
        results.update(parsed)
        errors.update(errs)

    source = Path(state["source_workbook"])
    wb = load_workbook(source)
    try:
        ws = wb.active
        type_col = find_or_add_col(ws, "event_type")
        certainty_col = find_or_add_col(ws, "event_type_certainty")
        review_col = find_or_add_col(ws, "event_type_needs_review")
        reason_col = find_or_add_col(ws, "event_type_reason")

        prev_col = find_or_add_col(ws, "event_type_prev")
        for cid, job in state["jobs"].items():
            row = job["excel_row"]
            result = results.get(cid)
            if not result:
                continue
            old = ws.cell(row, type_col).value
            if old and old != result["event_type"]:
                ws.cell(row, prev_col, old)       # kept for review; blank when unchanged
            ws.cell(row, type_col, result["event_type"])
            ws.cell(row, certainty_col, result["certainty"])
            ws.cell(row, review_col, result["needs_review"])
            ws.cell(row, reason_col, result["reason"])

        output = source.parent / (source.stem + " - event types" + source.suffix)
        wb.save(output)
    finally:
        wb.close()

    print(f"\nFinished. Valid results: {len(results):,}")
    print(f"Errors: {len(errors):,}")
    print(f"Output:\n{output}")

def main():
    if len(sys.argv) < 2:
        raise SystemExit(
            f'Usage:\n'
            f'  python3 {Path(__file__).name} submit "table.xlsx"\n'
            f'  python3 {Path(__file__).name} status\n'
            f'  python3 {Path(__file__).name} retry_failed\n'
            f'  python3 {Path(__file__).name} download'
        )
    cmd = sys.argv[1].lower()
    if cmd == "submit":
        if len(sys.argv) != 3:
            raise SystemExit("submit requires an .xlsx filename")
        command_submit(sys.argv[2])
    elif cmd == "status":
        command_status()
    elif cmd == "retry_failed":
        command_retry_failed()
    elif cmd == "download":
        command_download()
    else:
        raise SystemExit(f"Unknown command: {cmd}")

if __name__ == "__main__":
    main()
