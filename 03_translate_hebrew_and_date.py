#!/usr/bin/env python3
# Both Sides project
# Requirements:
#   pip install openai openpyxl
#   export OPENAI_API_KEY="..."

import json
import os
import sys
from pathlib import Path
from openai import OpenAI
from openpyxl import load_workbook

MODEL = "gpt-5.6-terra"
REASONING_EFFORT = "medium"
CHUNK_SIZE = 250
MAX_ACTIVE_BATCHES = 2

STATE_FILE = "date_hebrew_state.json"
BATCH_DIR = "date_hebrew_batches"

INSTRUCTIONS = r"""
You are translating and editing an English political-event description into concise, natural Hebrew for a data-journalism table.

You receive:
- MAIN ACTOR — context only.
- EVENT TYPE — context only.
- ENGLISH DESCRIPTION — the factual source text, with the leading event date already removed.

The English description is authoritative.
MAIN ACTOR and EVENT TYPE may help you understand the event, but they must NEVER cause you to add a fact that is not in the English description.


OUTPUT GOAL

Produce `description_he_medium`: a concise Hebrew factual summary of the event.

The Hebrew should read like a professionally edited table entry, not like a literal machine translation.


MANDATORY STYLE

- Write in natural, fluent Hebrew.
- Use PAST TENSE.
- Use correct Hebrew grammar, spelling, agreement, numbers, and punctuation.
- Do NOT include the event date.
- Do NOT begin with a date or Hebrew month.
- Do NOT mention the source/outlet unless it is substantively part of the event.
- Do NOT explain, interpret, editorialize, or add context.
- Do NOT use present tense to narrate the event.
- Do NOT invent motives, identities, outcomes, or causal relationships.


WHAT TO PRESERVE

Preserve the central factual information:
- who acted
- what they did
- where it happened
- meaningful participant numbers
- injuries, arrests, detentions, or other important consequences
- the stated political demand/motive when it is central to understanding the action
- specific place names when useful

Specific local place names may remain when they add useful information.
For example, do not automatically delete a place such as Kfar Tapuach simply because it is specific.

If the English text is long, preserve only the details that materially distinguish the event.


WHAT TO REMOVE OR COMPRESS

Usually remove or compress:
- the opening date
- repetitive wording
- administrative geographic parentheticals
- redundant descriptions of the same action
- long lists of organizations when they are not essential
- long lists of speakers or participants unless they materially distinguish the event
- explanatory background already obvious from the event
- source-like hedging or repetitive reporting language

Do not mechanically preserve every sentence.


LENGTH / COMPRESSION

The Hebrew MUST become progressively more compressed as the English source becomes longer.

Use these approximate targets:

SHORT source (up to ~250 English characters):
- usually about 60–130 Hebrew characters
- preserve most useful concrete details

MEDIUM source (~251–500 characters):
- usually about 100–190 Hebrew characters
- summarize; do not translate sentence by sentence

LONG source (~501–900 characters):
- usually about 140–260 Hebrew characters
- keep the central action plus only the most useful secondary details

VERY LONG source (over ~900 characters):
- usually about 180–320 Hebrew characters
- strongly compress; do not reproduce the structure of the original

These are guides, not quotas.
Clarity is more important than hitting an exact character count, but the Hebrew should clearly be shorter than the English.


EXAMPLES OF THE DESIRED STYLE

English:
"On 22 April 2023, at least 10,000 to about 28,000 protested in Netanya against the Netanyahu-led coalition's proposed judicial overhaul legislation. Former head of the Shin Bet Carmi Gillon spoke at the protest. Dozens of women dressed as 'handmaids' from Building an Alternative were present at the protest, in addition to health workers."

Good Hebrew:
"בין 10,000 ל־28,000 מפגינים מחו בנתניה נגד הרפורמה המשפטית. בהפגנה השתתפו גם נשות מחאת השפחות, אנשי צוות רפואי, וראש השב״כ לשעבר כרמי גילון שנאם במקום."

English:
"On 27 May 2023, about 76,000 to 135,000 protested in Tel Aviv city against the Netanyahu-led coalition's proposed judicial overhaul legislation. Former Likud MK and Defense Minister Moshe Yaalon spoke at the protest, in addition to protest leader and Black Flag Movement founder Shikma Bressler. LGBTQ activists, students, workers from the high-tech industry, female activists from Building an Alternative, Black Flag Movement activists, and anti-occupation activists all took part in the protest."

Good Hebrew:
"בין 76,000 ל־135,000 מפגינים מחו בתל אביב נגד הרפורמה המשפטית. בין המשתתפים היו פעילי להט״ב, סטודנטים, עובדי הייטק, פעילות בונות אלטרנטיבה, פעילי הדגלים השחורים ופעילים נגד הכיבוש. משה יעלון ושקמה ברסלר נאמו בהפגנה."


FINAL COPY-EDIT

Before returning:
1. Check that there is no date at the beginning.
2. Check that the narration is in past tense.
3. Check Hebrew gender/number agreement.
4. Remove unnecessary repetition.
5. Make sure the Hebrew is shorter than the English source.
6. Make sure no factual detail was invented.

"""

SCHEMA = {'type': 'object', 'properties': {'description_he_medium': {'type': 'string'}}, 'required': ['description_he_medium'], 'additionalProperties': False}


import re
from datetime import datetime

MONTHS = {
    "january":1, "february":2, "march":3, "april":4,
    "may":5, "june":6, "july":7, "august":8,
    "september":9, "october":10, "november":11, "december":12,
}

def extract_leading_date(description):
    text = description.strip()
    month_names = "|".join(MONTHS.keys())
    m = re.search(
        rf"(?i)\b(\d{{1,2}})(?:st|nd|rd|th)?\s+({month_names})\s+(20\d{{2}})\b",
        text[:220],
    )
    if not m:
        return None
    try:
        return datetime(
            int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1))
        )
    except ValueError:
        return None

def remove_leading_event_date(description):
    text = description.strip()
    month = (
        r"(?:January|February|March|April|May|June|"
        r"July|August|September|October|November|December)"
    )
    patterns = [
        rf"(?i)^\s*(?:On|Around|About|Approximately|During|Early on|Late on)?\s*"
        rf"\d{{1,2}}(?:st|nd|rd|th)?\s*[-–]\s*\d{{1,2}}(?:st|nd|rd|th)?\s+"
        rf"{month}\s+20\d{{2}}\s*[,;:.-]?\s*",
        rf"(?i)^\s*(?:On|Around|About|Approximately|During|Early on|Late on)?\s*"
        rf"\d{{1,2}}(?:st|nd|rd|th)?\s+{month}\s+20\d{{2}}\s*[,;:.-]?\s*",
        rf"(?i)^\s*(?:Between|From)\s+\d{{1,2}}(?:st|nd|rd|th)?\s+"
        rf"(?:and|to|-|–)\s+\d{{1,2}}(?:st|nd|rd|th)?\s+"
        rf"{month}\s+20\d{{2}}\s*[,;:.-]?\s*",
    ]
    for p in patterns:
        cleaned = re.sub(p, "", text, count=1).strip()
        if cleaned != text:
            return cleaned
    return text

def length_bin(text):
    n = len(text)
    if n <= 250:
        return "short"
    if n <= 500:
        return "medium"
    if n <= 900:
        return "long"
    return "very_long"

def strip_leading_hebrew_date(summary):
    months = (
        r"ינואר|פברואר|מרץ|אפריל|מאי|יוני|יולי|"
        r"אוגוסט|ספטמבר|אוקטובר|נובמבר|דצמבר"
    )
    patterns = [
        rf"^\s*(?:ביום\s+)?(?:ב[־-]?)?\d{{1,2}}\s+"
        rf"(?:ב[־-]?)?(?:{months})(?:\s+20\d{{2}})?\s*[,;:.\-–—]?\s*",
        rf"^\s*(?:בחודש\s+)?(?:ב[־-]?)?(?:{months})(?:\s+20\d{{2}})?"
        rf"\s*[,;:.\-–—]?\s*",
        r"^\s*(?:בשנת\s+|ב[־-]?)?20\d{2}\s*[,;:.\-–—]?\s*",
    ]
    cleaned = summary.strip()
    for _ in range(2):
        before = cleaned
        for p in patterns:
            candidate = re.sub(p, "", cleaned, count=1).strip()
            if candidate != cleaned:
                cleaned = candidate
                break
        if cleaned == before:
            break
    return cleaned


def norm(v):
    return "" if v is None else str(v).strip()

def hnorm(v):
    return norm(v).lower().replace(" ", "_")

def find_col(headers, names, required=False):
    lookup = {hnorm(v): i + 1 for i, v in enumerate(headers) if v is not None}
    for name in names:
        c = lookup.get(hnorm(name))
        if c:
            return c
    if required:
        raise KeyError(f"Missing required column. Tried: {names}")
    return None

def find_or_add_col(ws, name):
    headers = [c.value for c in ws[1]]
    c = find_col(headers, [name], required=False)
    if c:
        return c
    c = ws.max_column + 1
    ws.cell(1, c, name)
    return c

def extract_output_text(body):
    for item in body.get("output", []):
        if item.get("type") == "message":
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    return content.get("text", "")
    return ""

def save_state(base, state):
    (base / STATE_FILE).write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

def load_state(base):
    p = base / STATE_FILE
    if not p.exists():
        raise FileNotFoundError(f"{STATE_FILE} not found. Run submit first.")
    return json.loads(p.read_text(encoding="utf-8"))

def request_body(user_text):
    return {
        "model": MODEL,
        "reasoning": {"effort": REASONING_EFFORT},
        "instructions": INSTRUCTIONS,
        "input": user_text,
        "text": {
            "format": {
                "type": "json_schema",
                "name": "date_hebrew_result",
                "strict": True,
                "schema": SCHEMA,
            }
        },
        "max_output_tokens": 700,
    }

def submit_chunk(client, base, chunk):
    p = Path(chunk["input_path"])
    with p.open("rb") as f:
        uploaded = client.files.create(file=f, purpose="batch")
    batch = client.batches.create(
        input_file_id=uploaded.id,
        endpoint="/v1/responses",
        completion_window="24h",
    )
    chunk["input_file_id"] = uploaded.id
    chunk["batch_id"] = batch.id
    chunk["status"] = batch.status
    chunk["attempts"] = chunk.get("attempts", 0) + 1
    print(f"[{chunk['chunk_number']:03d}] submitted {chunk['request_count']:,} -> {batch.id} ({batch.status})")

def refresh(client, state):
    for chunk in state["chunks"]:
        if not chunk.get("batch_id"):
            chunk["status"] = "pending"
            continue
        b = client.batches.retrieve(chunk["batch_id"])
        chunk["status"] = b.status
        counts = getattr(b, "request_counts", None)
        if counts:
            chunk["complete"] = getattr(counts, "completed", 0) or 0
            chunk["failed"] = getattr(counts, "failed", 0) or 0
            chunk["total"] = getattr(counts, "total", 0) or 0
        if getattr(b, "output_file_id", None):
            chunk["output_file_id"] = b.output_file_id
        if getattr(b, "error_file_id", None):
            chunk["error_file_id"] = b.error_file_id
        if b.status == "failed":
            chunk["batch_errors"] = str(getattr(b, "errors", "") or "")

def active_count(state):
    return sum(
        1 for c in state["chunks"]
        if c.get("status") in {"validating", "in_progress", "finalizing", "cancelling"}
    )

def submit_pending(client, base, state):
    slots = max(0, MAX_ACTIVE_BATCHES - active_count(state))
    n = 0
    for chunk in state["chunks"]:
        if n >= slots:
            break
        if chunk.get("status", "pending") != "pending":
            continue
        submit_chunk(client, base, chunk)
        n += 1
    return n

def build_jobs(ws):
    headers = [c.value for c in ws[1]]
    row_id_col = find_col(headers, ["row_id", "row id"], required=False)
    actor_col = find_col(headers, ["main_actor", "main actor"], required=False)
    event_col = find_col(headers, ["event_type", "event type"], required=False)
    desc_col = find_col(headers, ["Description", "description"], required=True)
    he_col = find_col(headers, ["description_he_medium"], required=False)
    recl_col = find_col(headers, ["reclassify"], required=False)

    jobs = []
    for excel_row in range(2, ws.max_row + 1):
        row_id = norm(ws.cell(excel_row, row_id_col).value) if row_id_col else ""
        actor = norm(ws.cell(excel_row, actor_col).value) if actor_col else ""
        event_type = norm(ws.cell(excel_row, event_col).value) if event_col else ""
        description = norm(ws.cell(excel_row, desc_col).value)
        if not description:
            continue
        # ROW SELECTION: translate only rows with no Hebrew yet, or reclassify = yes
        # (set by 04 after a split; 03 is the last step, so download clears it).
        recl = norm(ws.cell(excel_row, recl_col).value).lower() == "yes" if recl_col else False
        if not recl and he_col and norm(ws.cell(excel_row, he_col).value):
            continue
        clean = remove_leading_event_date(description)
        dt = extract_leading_date(description)
        cid = row_id or f"excel-row-{excel_row}"
        jobs.append({
            "custom_id": cid,
            "excel_row": excel_row,
            "row_id": row_id,
            "date_iso": dt.date().isoformat() if dt else "",
            "source_length": len(description),
            "length_bin": length_bin(description),
            "user_text": (
                f"MAIN ACTOR:\n{actor or '(not supplied)'}\n\n"
                f"EVENT TYPE:\n{event_type or '(not supplied)'}\n\n"
                f"SOURCE LENGTH CLASS:\n{length_bin(description)}\n\n"
                f"ENGLISH DESCRIPTION:\n{clean}"
            ),
        })

    return jobs

def command_submit(filename):
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    source = Path(filename).expanduser().resolve()
    base = source.parent
    batch_dir = base / BATCH_DIR
    batch_dir.mkdir(exist_ok=True)

    wb = load_workbook(source)  # full load: ws.cell() is O(1); read_only made it O(rows) per call
    try:
        ws = wb.active
        jobs = build_jobs(ws)
    finally:
        wb.close()

    chunks = []
    for start in range(0, len(jobs), CHUNK_SIZE):
        part = jobs[start:start + CHUNK_SIZE]
        number = len(chunks) + 1
        jsonl = batch_dir / f"date_hebrew_{number:03d}.jsonl"
        with jsonl.open("w", encoding="utf-8") as f:
            for job in part:
                req = {
                    "custom_id": job["custom_id"],
                    "method": "POST",
                    "url": "/v1/responses",
                    "body": request_body(job["user_text"]),
                }
                f.write(json.dumps(req, ensure_ascii=False) + "\n")
        chunks.append({
            "chunk_number": number,
            "input_path": str(jsonl),
            "request_count": len(part),
            "custom_ids": [j["custom_id"] for j in part],
            "status": "pending",
            "batch_id": None,
            "attempts": 0,
        })

    state = {
        "source_workbook": str(source),
        "rows_submitted": len(jobs),
        "jobs": {j["custom_id"]: j for j in jobs},
        "chunks": chunks,
    }
    save_state(base, state)

    client = OpenAI()
    submit_pending(client, base, state)
    save_state(base, state)

    print(f"Prepared {len(jobs):,} rows in {len(chunks)} chunks.")
    print("Run periodically:")
    print(f"python3 {Path(__file__).name} status")

def command_status():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    newly = submit_pending(client, base, state)
    save_state(base, state)

    print(f"\nChecking {len(state['chunks'])} chunks...\n")
    for c in state["chunks"]:
        total = c.get("total", c["request_count"])
        print(
            f"[{c['chunk_number']:03d}] {c.get('status','pending'):<12} "
            f"{c.get('complete',0):,}/{total:,} complete, "
            f"{c.get('failed',0):,} failed"
        )
        if c.get("status") == "failed" and c.get("batch_errors"):
            print("   ", c["batch_errors"][:500])

    completed = sum(c.get("status") == "completed" for c in state["chunks"])
    failed = sum(c.get("status") == "failed" for c in state["chunks"])
    pending = sum(c.get("status") == "pending" for c in state["chunks"])

    print(f"\nCompleted chunks: {completed}/{len(state['chunks'])}")
    print(f"Pending chunks: {pending}")
    print(f"Failed chunks: {failed}")
    if newly:
        print(f"Submitted {newly} new chunk(s).")
    if failed:
        print(f"Run: python3 {Path(__file__).name} retry_failed")
    if completed == len(state["chunks"]):
        print(f"Run: python3 {Path(__file__).name} download")

def command_retry_failed():
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    reset = 0
    for c in state["chunks"]:
        if c.get("status") == "failed":
            c["batch_id"] = None
            c["status"] = "pending"
            c["input_file_id"] = None
            c["output_file_id"] = None
            c["error_file_id"] = None
            reset += 1
    submit_pending(client, base, state)
    save_state(base, state)
    print(f"Reset {reset} failed chunk(s).")

def parse_output(content):
    results, errors = {}, {}
    for line in content.decode("utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        cid = obj.get("custom_id")
        if obj.get("error"):
            errors[cid] = str(obj["error"])
            continue
        response = obj.get("response") or {}
        if response.get("status_code") != 200:
            errors[cid] = f"HTTP {response.get('status_code')}"
            continue
        text = extract_output_text(response.get("body") or {})
        if not text:
            errors[cid] = "No output_text"
            continue
        try:
            results[cid] = json.loads(text)
        except Exception as exc:
            errors[cid] = f"JSON parse error: {exc}"
    return results, errors

def command_download():
    if not os.getenv("OPENAI_API_KEY"):
        raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd()
    state = load_state(base)
    client = OpenAI()
    refresh(client, state)
    save_state(base, state)

    incomplete = [c for c in state["chunks"] if c.get("status") != "completed"]
    if incomplete:
        raise RuntimeError("Not all chunks are completed.")

    results, errors = {}, {}
    for i, c in enumerate(state["chunks"], 1):
        b = client.batches.retrieve(c["batch_id"])
        print(f"[{i:03d}/{len(state['chunks']):03d}] downloading...")
        content = client.files.content(b.output_file_id)
        parsed, errs = parse_output(content.content)
        results.update(parsed)
        errors.update(errs)

    source = Path(state["source_workbook"])
    wb = load_workbook(source)
    try:
        ws = wb.active
        date_col = find_or_add_col(ws, "date")
        he_col = find_or_add_col(ws, "description_he_medium")
        recl_col = find_or_add_col(ws, "reclassify")

        for cid, job in state["jobs"].items():
            row = job["excel_row"]
            result = results.get(cid)
            if not result:
                continue
            if job.get("date_iso"):
                dt = datetime.strptime(job["date_iso"], "%Y-%m-%d")
                ws.cell(row, date_col, dt)
                ws.cell(row, date_col).number_format = "DD/MM/YYYY"
            summary = strip_leading_hebrew_date(result["description_he_medium"])
            ws.cell(row, he_col, summary)
            ws.cell(row, recl_col, None)

        output = source.parent / (source.stem + " - date hebrew" + source.suffix)
        wb.save(output)
    finally:
        wb.close()

    print(f"\nFinished. Valid results: {len(results):,}")
    print(f"Errors: {len(errors):,}")
    print(f"Output:\n{output}")

def main():
    if len(sys.argv) < 2:
        raise SystemExit(
            f'Usage:\n'
            f'  python3 {Path(__file__).name} submit "table.xlsx"\n'
            f'  python3 {Path(__file__).name} status\n'
            f'  python3 {Path(__file__).name} retry_failed\n'
            f'  python3 {Path(__file__).name} download'
        )
    cmd = sys.argv[1].lower()
    if cmd == "submit":
        if len(sys.argv) != 3:
            raise SystemExit("submit requires an .xlsx filename")
        command_submit(sys.argv[2])
    elif cmd == "status":
        command_status()
    elif cmd == "retry_failed":
        command_retry_failed()
    elif cmd == "download":
        command_download()
    else:
        raise SystemExit(f"Unknown command: {cmd}")

if __name__ == "__main__":
    main()
