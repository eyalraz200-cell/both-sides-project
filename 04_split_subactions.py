#!/usr/bin/env python3
# Both Sides project — step 04: split one ACLED row into one event per action.
# Reuses the Batch pipeline from 02 (import by file name) and only changes the
# prompt, the schema, which rows are sent, and what `download` writes back.
#
#   python3 04_split_subactions.py submit events.xlsx
#   python3 04_split_subactions.py status | retry_failed | download
#
# download rewrites the parent row (Description, crowd; original kept in
# description_original / crowd_total), APPENDS one child row per extra action
# (new row_id, split_from = parent), blanks event_type / description_he_medium
# on every touched row and sets reclassify = yes so 02 and 03 pick them up.
# Rows with a date in event_type_hand or description_hand are never sent, so a
# split never undoes a hand decision.

import importlib.util, json, os, re, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("step02", HERE / "02_classify_event_type_no_intimidation.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)

m.STATE_FILE = "split_state.json"
m.BATCH_DIR = "split_batches"

m.INSTRUCTIONS = r"""
You split a political-event record into ONE EVENT PER DISTINCT ACTION when the record
describes a subset of the MAIN ACTOR doing something separate from the main crowd.

You receive: MAIN ACTOR, CURRENT EVENT TYPE, CROWD (the whole-event figure, may be blank),
and the ENGLISH DESCRIPTION (starts with the date).

WHEN TO SPLIT
Split when the description contains BOTH:
(a) a main action by the whole crowd (a rally, march, demonstration, raid), AND
(b) a different action by a SUBSET of that same MAIN ACTOR, marked by its own size
    ("dozens", "hundreds", "about 200", "some protesters", "a group of", "several",
    "a number of", "part of the crowd") or by a different time/place the same day
    ("following the main demonstration", "later", "toward the end", "separately").
Typical: "about 40,000 protested in Tel Aviv ... dozens blocked the Ayalon Highway."
         "hundreds marched ... several protesters set tires on fire."
         "hundreds of settlers demonstrated ... some of them threw stones at cars."
Also split when two clearly separate actions by the same actor at different places are
packed into one record ("protests took place in X and in Y ...") ONLY if each has its own
size or location detail; otherwise keep one event.

WHEN NOT TO SPLIT
- The whole crowd did it ("thousands blocked traffic", "hundreds clashed with police").
- The extra action is by a SECONDARY ACTOR: police, soldiers, passing drivers, passers-by,
  counter-protesters, Palestinians repelling settlers, "a rioter" who attacked the
  protesters. Never create an event for them.
- Background sentences (what the protest was about, who spoke, context of the week).
- Only one action is described, however long the text.

OUTPUT
`split`: "yes" or "no".
`events`: 1 to 3 events. If split = "no", exactly one event: the original description
unchanged (copy it verbatim) and crowd_text = CROWD as given (or "" if blank).
If split = "yes": the FIRST event is the main action with the whole-event crowd
(crowd_text = CROWD if given, else the main figure in the text); each further event is
one subset action.

For every event:
- `description_en`: a self-contained English description in ACLED's register, starting
  "On <date>, ..." exactly as the original does. Keep ACLED's own wording and facts;
  move sentences, do not invent. The main event keeps the context (what the protest was
  about, organisers, speakers). The subset event gets a short context clause so it stands
  alone ("Following the main anti-overhaul demonstration in Tel Aviv, dozens blocked
  traffic on Ayalon Highway ..."). Do not repeat the subset action inside the main event.
  Police response belongs to the event it responded to.
- `crowd_text`: the size of THIS event's group, in the source's own words: "about 40,000",
  "dozens", "hundreds", "about 200". A SUBSET event takes ONLY a figure that the text
  attaches to the subset itself. If the subset is "some protesters", "demonstrators",
  "a group", "several", "a women's group", "activists", "part of the crowd" — anything
  without its own number — write "unspecified". NEVER copy the main event's figure or the
  CROWD field into a subset event. For the main event with a blank CROWD and no figure in
  the text, write "".
- `is_primary`: true for the main action, false for every subset event.

Never change the MAIN ACTOR. Never classify the event type — that is done later.
Do not add, soften or sharpen any fact.
"""

m.SCHEMA = {
    "type": "object",
    "properties": {
        "split": {"type": "string", "enum": ["yes", "no"]},
        "reason": {"type": "string"},
        "events": {
            "type": "array",
            "minItems": 1, "maxItems": 3,
            "items": {
                "type": "object",
                "properties": {
                    "is_primary": {"type": "boolean"},
                    "crowd_text": {"type": "string"},
                    "description_en": {"type": "string"},
                },
                "required": ["is_primary", "crowd_text", "description_en"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["split", "reason", "events"],
    "additionalProperties": False,
}

_orig_request_body = m.request_body
def request_body(user_text):
    b = _orig_request_body(user_text)
    b["text"]["format"]["name"] = "split_result"
    b["max_output_tokens"] = 2500
    return b
m.request_body = request_body

# Only rows that can plausibly hold a second action: two different crowd figures, or
# subset wording. Everything else is left untouched (and costs nothing).
SUB = re.compile(r"\b(some|a group|a number|several|a few|dozens|hundreds|tens|part|\d+)\s+of (the )?(protesters|demonstrators|activists|rioters|them|settlers|participants|marchers)|following the (main )?(protest|demonstration|rally)|after the (main )?(protest|demonstration|rally)|toward the end of the (protest|demonstration)|(separately|meanwhile|later),? (dozens|hundreds|some|a group|several)|(dozens|hundreds|some|several) (then |also |later )?(blocked|marched|clashed|set fire|lit|broke)", re.I)
NUM = re.compile(r"\b(dozens|hundreds|thousands|tens of thousands|about [\d,]+|around [\d,]+|at least [\d,]+|over [\d,]+)\b", re.I)
ACT = re.compile(r"\b(blocked|block traffic|clashed|broke through|set (fire|tires)|lit (a |)(bonfire|fire|tires)|burned tires|chained|stormed|broke into|threw)\b", re.I)
SOLE_SOURCE_EXCLUDE = {"plo negotiations affairs department"}   # same set as server.py: never shipped

def build_jobs(ws):
    headers = [c.value for c in ws[1]]
    col = lambda *names, req=False: m.find_col(headers, list(names), required=req)
    c_id, c_actor, c_desc = col("row_id"), col("main_actor", req=True), col("Description", req=True)
    c_type, c_crowd, c_src, c_hid, c_ds = col("event_type"), col("crowd"), col("source"), col("hidden"), col("data_source")
    c_split = col("split_from")
    c_corr = col("corroborated_by")
    c_thand, c_dhand = col("event_type_hand"), col("description_hand")
    jobs = []
    for r in range(2, ws.max_row + 1):
        g = lambda c: m.norm(ws.cell(r, c).value) if c else ""
        if g(c_ds) not in ("acled", "manual"): continue           # Fortress rows are single-action Hebrew lines
        if g(c_hid) or g(c_split): continue
        actor, desc = g(c_actor), g(c_desc)
        if not actor or actor == "not relevant" or not desc: continue
        srcs = {s.strip().lower() for s in g(c_src).split(";") if s.strip()}
        if srcs and srcs <= SOLE_SOURCE_EXCLUDE and not g(c_corr): continue   # corroborated rows ship
        # A hand-set type or Hebrew line is never undone: download would blank both and
        # send the row back through 02/03 (CLAUDE.md: never re-run 02 over a row with a
        # date in event_type_hand).
        if g(c_thand) or g(c_dhand): continue
        # Second-pass gate (2026-10-03): any protest row above הפגנה לא אלימה with a crowd
        # figure and a sub-action verb, plus the original subset-wording / two-figures test.
        above = g(c_type) not in ("", "הפגנה לא אלימה")
        if not (SUB.search(desc) or len({x.lower() for x in NUM.findall(desc)}) >= 2
                or (above and NUM.search(desc) and ACT.search(desc))): continue
        if g(c_split) or (col("description_original") and g(col("description_original"))): continue   # already split once
        jobs.append({
            "custom_id": g(c_id) or f"excel-row-{r}", "excel_row": r, "row_id": g(c_id),
            "user_text": (f"MAIN ACTOR:\n{actor}\n\nCURRENT EVENT TYPE:\n{g(c_type) or '(blank)'}\n\n"
                          f"CROWD:\n{g(c_crowd) or '(blank)'}\n\nENGLISH DESCRIPTION:\n{desc}"),
        })
    return jobs
m.build_jobs = build_jobs

def command_download():
    from openai import OpenAI
    from openpyxl import load_workbook
    if not os.getenv("OPENAI_API_KEY"): raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd(); state = m.load_state(base); client = OpenAI()
    m.refresh(client, state); m.save_state(base, state)
    if any(c.get("status") != "completed" for c in state["chunks"]): raise RuntimeError("Not all chunks are completed.")
    results, errors = {}, {}
    for c in state["chunks"]:
        b = client.batches.retrieve(c["batch_id"])
        parsed, errs = m.parse_output(client.files.content(b.output_file_id).content)
        results.update(parsed); errors.update(errs)

    source = Path(state["source_workbook"]); wb = load_workbook(source); ws = wb.active
    try:
        H = lambda name: m.find_or_add_col(ws, name)
        c = {n: H(n) for n in ["row_id", "Description", "crowd", "event_type", "description_he_medium",
                               "description_original", "crowd_total", "split_from", "reclassify", "split_reason",
                               "event_type_hand", "main_actor_hand"]}
        headers = [x.value for x in ws[1]]
        ids = [int(str(ws.cell(r, c["row_id"]).value).split("-")[1]) for r in range(2, ws.max_row + 1) if ws.cell(r, c["row_id"]).value]
        next_id = max(ids) + 1
        parents = children = 0
        for cid, job in state["jobs"].items():
            res = results.get(cid)
            if not res or res.get("split") != "yes" or len(res["events"]) < 2: continue
            r = job["excel_row"]
            evs = sorted(res["events"], key=lambda e: not e["is_primary"])
            primary, subs = evs[0], evs[1:]
            if not ws.cell(r, c["description_original"]).value:
                ws.cell(r, c["description_original"], ws.cell(r, c["Description"]).value)
                ws.cell(r, c["crowd_total"], ws.cell(r, c["crowd"]).value)
            ws.cell(r, c["Description"], primary["description_en"])
            ws.cell(r, c["crowd"], primary["crowd_text"] or None)
            setattr(ws.cell(r, c["event_type"]), 'value', None); setattr(ws.cell(r, c["description_he_medium"]), 'value', None)
            setattr(ws.cell(r, c["event_type_hand"]), 'value', None); ws.cell(r, c["reclassify"], "yes"); ws.cell(r, c["split_reason"], res.get("reason"))
            parents += 1
            for s in subs:
                nr = ws.max_row + 1
                for col_i in range(1, len(headers) + 1):            # copy every field, then override
                    ws.cell(nr, col_i, ws.cell(r, col_i).value)
                ws.cell(nr, c["row_id"], f"row-{next_id}"); next_id += 1
                ws.cell(nr, c["Description"], s["description_en"])
                ws.cell(nr, c["crowd"], None if s["crowd_text"] in ("", "unspecified") else s["crowd_text"])
                setattr(ws.cell(nr, c["crowd_total"]), 'value', None); setattr(ws.cell(nr, c["description_original"]), 'value', None)
                ws.cell(nr, c["split_from"], job["row_id"]); ws.cell(nr, c["reclassify"], "yes")
                for k in ("event_type", "description_he_medium", "event_type_hand", "main_actor_hand", "split_reason"):
                    setattr(ws.cell(nr, c[k]), 'value', None)
                children += 1
        out = source.parent / (source.stem + " - split" + source.suffix); wb.save(out)
    finally:
        wb.close()
    print(f"Split {parents:,} rows into {parents + children:,} events ({children:,} new rows). Errors: {len(errors):,}.\nOutput:\n{out}")
    print("Next: python3 02_classify_event_type_no_intimidation.py submit <that file>  (rows with reclassify=yes), then 03.")

def main():
    if len(sys.argv) < 2:
        raise SystemExit(f"Usage: python3 {Path(__file__).name} submit events.xlsx | status | retry_failed | download")
    cmd = sys.argv[1].lower()
    if cmd == "submit": m.command_submit(sys.argv[2])
    elif cmd == "status": m.command_status()
    elif cmd == "retry_failed": m.command_retry_failed()
    elif cmd == "download": command_download()
    else: raise SystemExit(f"Unknown command: {cmd}")

if __name__ == "__main__":
    main()
