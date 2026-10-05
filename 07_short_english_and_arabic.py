#!/usr/bin/env python3
# Both Sides project — step 07: the shipped English line (short) and the Arabic line.
# Reuses the Batch pipeline from 02 (import by file name).
#
#   python3 07_short_english_and_arabic.py submit events.xlsx
#   python3 07_short_english_and_arabic.py status | retry_failed | download
#
# Rows sent: live rows (not hidden, actor not "not relevant") with a Hebrew line and
# a blank `description_ar`, EXCEPT Fortress rows not yet rewritten by step 06
# (blank `description_he_original`). Rows with a date in `description_hand` skip.
# download writes `description_en_short` (Description itself — ACLED's full English,
# which the classifiers read — is never touched) and `description_ar`.
# server.py ships description_en_short (falling back to Description) in
# events-en.json and description_ar in events-ar.json.

import importlib.util, os
from pathlib import Path

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("step02", HERE / "02_classify_event_type_no_intimidation.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)

m.STATE_FILE = "short_en_ar_state.json"
m.BATCH_DIR = "short_en_ar_batches"
m.REASONING_EFFORT = "low"
m.MAX_ACTIVE_BATCHES = 8          # big run; each batch is small (250 short calls)

m.INSTRUCTIONS = r"""
You write two lines for one event in a dataset of Israeli political events: a SHORT
ENGLISH line and an ARABIC line. Both must carry exactly the facts of the HEBREW line,
which is the reference text shown on the site.

You receive the HEBREW line and, for context only, the ORIGINAL ENGLISH notes (ACLED's,
often long, starting with the date) — or "(none)".

SHORT ENGLISH (`description_en`)
- A translation of the HEBREW line, fact for fact: who, what, where, how many, outcome.
  NOTHING that is not in the Hebrew — not even if the original English has it (background,
  quotes, police statements, organisers, "according to reports"). The original English is
  only a spelling reference.
- LENGTH: you are told the Hebrew line's length in characters. The English must be at most
  1.4 × that length, and NEVER over 360 characters (the longest Hebrew line is 342). Shorter is
  fine if no fact is lost; to fit, use tighter wording, never drop a fact.
- Never start with the date or with an ACLED tag ("Property destruction:"). Start with
  the actor: "Hundreds of anti-overhaul protesters ...", "Israeli settlers ...".
- Place and person names ALWAYS from the ORIGINAL ENGLISH when it names them — never guess a
  different village. Spellings and terms from the ORIGINAL ENGLISH where it has them ("Huwara", "Masafer Yatta",
  "Kaplan Street", "Ayalon Highway", "judicial overhaul", "Palestinian-owned").

ARABIC (`description_ar`)
- Modern Standard Arabic, neutral news register (as in Arab48 / Al Jazeera news copy).
- A faithful translation of the HEBREW line — same facts, same order. Do not add or
  soften anything. Length: at most 1.4 × the Hebrew line's length, never over 360 characters.
- Terms (always these): settlers = مستوطنون; Israeli settlers = مستوطنون إسرائيليون; the judicial
  overhaul / reform = التعديلات القضائية; Haredim = الحريديم; Haredi draft refusers (עריקים) =
  المتهربون من التجنيد; the hostages (החטופים) = المختطفون — add "في غزة" or "لدى حماس" ONLY when
  the Hebrew says so; Arab citizens of
  Israel (ערבים ישראלים) = المواطنون العرب في إسرائيل; the West Bank = الضفة الغربية; Jerusalem =
  القدس; the Temple Mount / al-Aqsa compound = المسجد الأقصى; Knesset = الكنيست.
- Detention/abduction verbs follow the Hebrew: החזיקו בכפייה = احتجزوا بالقوة; חטפו = اختطفوا;
  לקחו בכפייה = اقتادوا بالقوة.
- Place names: the standard Arabic name for Arab towns and villages (حوارة, ترمسعيا,
  مسافر يطا, أم الفحم); Israeli places in their usual Arabic transliteration (تل أبيب,
  شارع كابلان, هرتسليا).
- Numbers in Western digits (200, 10,000).
- NEVER add a word the Hebrew does not justify: no "Israeli", "in Gaza", "Palestinian" unless it is there.

Return both fields; `note` is "" unless the Hebrew is unclear or contradicts the English.
"""

m.SCHEMA = {
    "type": "object",
    "properties": {
        "description_en": {"type": "string"},
        "description_ar": {"type": "string"},
        "note": {"type": "string"},
    },
    "required": ["description_en", "description_ar", "note"],
    "additionalProperties": False,
}

_orig_request_body = m.request_body
def request_body(user_text):
    b = _orig_request_body(user_text)
    b["text"]["format"]["name"] = "short_en_ar"
    b["max_output_tokens"] = 2500
    return b
m.request_body = request_body


SOLE_SOURCE_EXCLUDE = {"plo negotiations affairs department"}   # same set as server.py: never shipped

def build_jobs(ws):
    headers = [c.value for c in ws[1]]
    col = lambda *names, req=False: m.find_col(headers, list(names), required=req)
    c_id, c_actor, c_he, c_en = col("row_id"), col("main_actor", req=True), col("description_he_medium", req=True), col("Description")
    c_ar, c_hid, c_ds, c_orig, c_hand = col("description_ar"), col("hidden"), col("data_source"), col("description_he_original"), col("description_hand")
    c_note = col("description_note")
    c_src = m.find_col(headers, ["source"], required=False)
    c_corr = m.find_col(headers, ["corroborated_by"], required=False)
    jobs = []
    for r in range(2, ws.max_row + 1):
        g = lambda c: m.norm(ws.cell(r, c).value) if c else ""
        srcs = {x.strip().lower() for x in g(c_src).split(";") if x.strip()}
        if srcs and srcs <= SOLE_SOURCE_EXCLUDE and not g(c_corr): continue   # corroborated rows ship
        actor = g(c_actor)
        if not actor or actor == "not relevant" or g(c_hid) or g(c_hand): continue
        if g(c_ar) or not g(c_he): continue
        if g(c_ds) == "the fortress" and not g(c_orig): continue     # wait for step 06
        en = g(c_en)        # ACLED: the full notes; Fortress: step 06's English (place spellings)
        hint = ""
        note = m.norm(ws.cell(r, c_note).value) if c_note else ""
        if note.startswith("too long"):
            hint = f"\n\nYOUR PREVIOUS ATTEMPT WAS TOO LONG ({note}). Stay under {min(360, int(len(g(c_he)) * 1.4))} characters in both languages."
        jobs.append({
            "custom_id": g(c_id) or f"excel-row-{r}", "excel_row": r, "row_id": g(c_id),
            "user_text": f"HEBREW ({len(g(c_he))} characters):\n{g(c_he)}\n\nORIGINAL ENGLISH (spelling reference only):\n{en or '(none)'}" + hint,
        })
    return jobs
m.build_jobs = build_jobs


def command_download():
    from openai import OpenAI
    from openpyxl import load_workbook
    if not os.getenv("OPENAI_API_KEY"): raise EnvironmentError("OPENAI_API_KEY is not set.")
    base = Path.cwd(); state = m.load_state(base); client = OpenAI()
    m.refresh(client, state); m.save_state(base, state)
    if any(c.get("status") != "completed" for c in state["chunks"]): raise RuntimeError("Not all chunks are completed.")
    results, errors = {}, {}
    for c in state["chunks"]:
        b = client.batches.retrieve(c["batch_id"])
        parsed, errs = m.parse_output(client.files.content(b.output_file_id).content)
        results.update(parsed); errors.update(errs)
    source = Path(state["source_workbook"]); wb = load_workbook(source); ws = wb.active
    try:
        H = lambda n: m.find_or_add_col(ws, n)
        c_short, c_ar, c_note, c_id = H("description_en_short"), H("description_ar"), H("description_note"), H("row_id")
        # map by row_id, not excel_row: other steps may have appended rows since submit
        rowmap = {ws.cell(r, c_id).value: r for r in range(2, ws.max_row + 1)}
        written = long_ = 0
        for cid, job in state["jobs"].items():
            res = results.get(cid); r = rowmap.get(job["row_id"])
            if not res or not r: continue
            en, ar = m.norm(res["description_en"]), m.norm(res["description_ar"])
            if not en or not ar: continue
            he_len = len(m.norm(job["user_text"].split("\n")[1]))
            cap = min(380, max(60, int(he_len * 1.6)))
            if len(en) > cap or len(ar) > cap:
                ws.cell(r, c_note, f"too long: en {len(en)}, ar {len(ar)}, hebrew {he_len}")
                long_ += 1; continue           # left blank: the next run of 07 re-sends it with this note
            ws.cell(r, c_short, m.xl_safe(en)); ws.cell(r, c_ar, m.xl_safe(ar))
            if res.get("note"): ws.cell(r, c_note, m.xl_safe(res["note"]))
            written += 1
        out = source.parent / (source.stem + " - short en ar" + source.suffix); wb.save(out)
    finally:
        wb.close()
    print(f"\nFinished. Written: {written:,}. Too long, left for a re-run: {long_:,}. Errors: {len(errors):,}\nOutput:\n{out}")
m.command_download = command_download

if __name__ == "__main__":
    m.main()
