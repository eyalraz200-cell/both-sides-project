#!/usr/bin/env python3
"""Dedupe new ACLED rows against the Fortress log (step 08; run right after the fetch).

  python3 08_dedupe_fortress.py events.xlsx            # judge + apply, in place (backup first)
  python3 08_dedupe_fortress.py events.xlsx --dry-run  # judge, print, write nothing

For every live ACLED row not yet checked (blank `dedupe_note`, not hidden, not PLO-only —
PLO-only rows never ship, so a Fortress twin of one simply stays live; decided 2026-10-07),
collect the live Fortress rows dated within ±1 day and ask the model whether one of them is
THE SAME INCIDENT (the rules are _dedupe/INSTRUCTIONS.md, the pass Claude readers did by hand
on 2026-10-03). ACLED wins:
  sure   → the Fortress row gets hidden = "dup of <acled row>", acled_duplicate_of, acled_match
           = sure, dedupe_note; the ACLED row's dedupe_note names the twin
  unsure → the Fortress row keeps shipping; acled_duplicate_of + acled_match = unsure + note,
           for the review
  none   → the ACLED row's dedupe_note = "no fortress twin"
An ACLED row with no Fortress candidate on those dates is stamped "no fortress candidates"
without a call. Synchronous Responses API calls, THREADS at a time (a week is ~100 calls).
"""
import datetime as dt
import json
import os
import shutil
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openai import OpenAI
from openpyxl import load_workbook

MODEL = "gpt-5.6-terra"
THREADS = 8
SOLE_SOURCE_EXCLUDE = {"plo negotiations affairs department"}
ARCHIVE = Path(__file__).resolve().parent / "_xlsx-archive"

INSTRUCTIONS = """You are deduplicating two event datasets about settler violence in the West Bank.
You get ONE ACLED event (English notes, usually starting with the date) and a list of
candidate rows from "The Fortress", a Hebrew settler-violence log, dated within ±1 day
(reporting dates differ by a day). Decide whether any candidate describes THE SAME INCIDENT:
same place, same action, same victims/outcome. Transliterate place names across
Hebrew/English yourself (ראס עין אל-עוג'א ≈ Al Awja / Ein al-Auja; קוסרה = Qusra; מוע'ייר =
Al Mughayyir; מסאפר יטא = Masafer Yatta / Yuta area; ח'ירבת = Khirbat; סינג'יל = Sinjil;
תורמוסעיא = Turmus Ayya; בורין = Burin; דומא = Duma; חווארה = Huwwara). Specific details
matter more than village names: counts of trees, sheep, injured, named people, a mosque, a
particular road. A village with several candidates on one day needs the details to pick one,
or none.
Verdicts: sure = clearly the same incident; unsure = compatible (same village, same kind of
act, same day) but details too thin to be certain; none = no candidate matches. If two
candidates both fit, pick the better one and name the other in the note. Prefer sure when
place and act line up and nothing contradicts. note: 10 words max, why."""

SCHEMA = {
    "type": "object", "additionalProperties": False,
    "properties": {
        "match": {"type": ["string", "null"]},
        "verdict": {"type": "string", "enum": ["sure", "unsure", "none"]},
        "note": {"type": "string"},
    },
    "required": ["match", "verdict", "note"],
}


def to_date(v):
    return v.date() if isinstance(v, dt.datetime) else dt.date.fromisoformat(str(v)[:10])


def judge(client, acled, cands):
    text = (f"ACLED event ({acled['date']}, {acled['loc']}, actor {acled['actor']}):\n{acled['desc']}\n\n"
            "Fortress candidates:\n" + "\n".join(
                f"- {c['row_id']} | {c['date']} | {c['loc']} | {c['cats']} | {c['desc']}" for c in cands))
    r = client.responses.create(
        model=MODEL, instructions=INSTRUCTIONS, input=text,
        reasoning={"effort": "medium"},
        text={"format": {"type": "json_schema", "name": "dedupe", "strict": True, "schema": SCHEMA}},
        max_output_tokens=2000,
    )
    out = json.loads(r.output_text)
    ids = {c["row_id"] for c in cands}
    if out["match"] not in ids:
        out["match"] = None
        if out["verdict"] != "none":
            out["verdict"], out["note"] = "none", "model named no valid candidate: " + out["note"]
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if not os.getenv("OPENAI_API_KEY"):
        sys.exit("OPENAI_API_KEY is not set.")
    src = Path(sys.argv[1]).resolve()
    dry = "--dry-run" in sys.argv
    today = str(dt.date.today())
    wb = load_workbook(src)
    ws = wb.active
    h = [c.value for c in ws[1]]
    c = {k: i + 1 for i, k in enumerate(h)}
    g = lambda r, k: ws.cell(r, c[k]).value

    fortress = {}       # date -> [cand dicts]
    todo = []
    for r in range(2, ws.max_row + 1):
        ds = g(r, "data_source")
        if g(r, "hidden") or not g(r, "date"):
            continue
        d = to_date(g(r, "date"))
        if ds == "the fortress":
            fortress.setdefault(d, []).append({
                "row": r, "row_id": g(r, "row_id"), "date": str(d), "loc": g(r, "location") or "",
                "cats": g(r, "fortress_categories") or "",
                "desc": g(r, "description_he_medium") or g(r, "description_he_original") or g(r, "Description") or "",
            })
        elif ds == "acled" and not g(r, "dedupe_note"):
            srcs = {s.strip().lower() for s in str(g(r, "source") or "").split(";") if s.strip()}
            if srcs and srcs <= SOLE_SOURCE_EXCLUDE:
                continue
            todo.append({"row": r, "row_id": g(r, "row_id"), "date": d, "loc": g(r, "location") or "",
                         "actor": g(r, "main_actor") or "", "desc": g(r, "Description") or ""})

    jobs, no_cand = [], []
    for a in todo:
        cands = [x for off in (-1, 0, 1) for x in fortress.get(a["date"] + dt.timedelta(days=off), [])]
        (jobs if cands else no_cand).append((a, cands))
    print(f"ACLED rows to check: {len(todo)}  — {len(jobs)} with Fortress candidates, {len(no_cand)} without")

    client = OpenAI()
    results = {}
    def work(item):
        a, cands = item
        try:
            return a["row_id"], judge(client, a, cands)
        except Exception as e:  # keep going; the row stays unchecked for the next run
            return a["row_id"], {"error": str(e)[:200]}
    with ThreadPoolExecutor(THREADS) as ex:
        for i, (rid, res) in enumerate(ex.map(work, jobs), 1):
            results[rid] = res
            if i % 50 == 0:
                print(f"  {i}/{len(jobs)}")

    sure = [(a, results[a["row_id"]]) for a, _ in jobs if results[a["row_id"]].get("verdict") == "sure"]
    unsure = [(a, results[a["row_id"]]) for a, _ in jobs if results[a["row_id"]].get("verdict") == "unsure"]
    errors = [a for a, _ in jobs if "error" in results[a["row_id"]]]
    print(f"sure {len(sure)}, unsure {len(unsure)}, none {len(jobs) - len(sure) - len(unsure) - len(errors)}, errors {len(errors)}")
    for a, res in unsure:
        print(f"  UNSURE {a['row_id']} ~ {res['match']}: {res['note']}")
    if dry:
        for a, res in sure[:20]:
            print(f"  sure {a['row_id']} = {res['match']}: {res['note']}")
        return

    frow = {x["row_id"]: x["row"] for lst in fortress.values() for x in lst}
    taken = set()
    for a, _ in no_cand:
        ws.cell(a["row"], c["dedupe_note"], f"no fortress candidates | {today}")
    for a, _ in jobs:
        res = results[a["row_id"]]
        if "error" in res:
            continue
        m = res["match"]
        if res["verdict"] == "sure" and m and m not in taken:
            taken.add(m)
            fr = frow[m]
            ws.cell(fr, c["hidden"], f"dup of {a['row_id']}")
            ws.cell(fr, c["acled_duplicate_of"], a["row_id"])
            ws.cell(fr, c["acled_match"], "sure")
            ws.cell(fr, c["dedupe_note"], f"{res['note']} | 08 {today}")
            ws.cell(a["row"], c["dedupe_note"], f"fortress twin {m} hidden | {today}")
        elif res["verdict"] == "unsure" and m:
            fr = frow[m]
            if not g(fr, "acled_duplicate_of"):
                ws.cell(fr, c["acled_duplicate_of"], a["row_id"])
                ws.cell(fr, c["acled_match"], "unsure")
                ws.cell(fr, c["dedupe_note"], f"{res['note']} | 08 {today}")
            ws.cell(a["row"], c["dedupe_note"], f"unsure twin {m} | {today}")
        else:
            ws.cell(a["row"], c["dedupe_note"], f"no fortress twin | {today}")
    ARCHIVE.mkdir(exist_ok=True)
    shutil.copy2(src, ARCHIVE / f"{src.stem}-before-dedupe-{today}{src.suffix}")
    wb.save(src)
    print(f"Applied: hid {len(taken)} Fortress rows. Saved {src.name}.")


if __name__ == "__main__":
    main()
